#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
钉钉群级通知适配器。

默认只 dry-run，不访问钉钉。只有显式传入 ``--send`` 才会向路由文件中的
HTTPS 群 Webhook 发消息。这里不实现企业负责人单聊，避免把群机器人能力
误包装成组织身份和单聊能力。
"""
import argparse
import base64
import datetime
import hashlib
import hmac
import json
import os
import re
import sys
import time
from urllib.parse import parse_qsl, urlencode, urlparse, urlunparse

import openpyxl
import requests
from notification_ledger import Ledger, key as delivery_key


NOTIFYABLE_CONCLUSIONS = {"匹配"}
SUCCESS_STATES = {"已推送", "已发送"}
TIMEOUT = 25


def _text(value):
    return "" if value is None else str(value).strip()


def norm_key(value):
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return _text(value)


def result_key(row):
    return f"{norm_key(row.get('政策ID', ''))}||{norm_key(row.get('企业名称', ''))}"


def read_xlsx(path):
    if not os.path.exists(path):
        return [], []
    workbook = openpyxl.load_workbook(path, data_only=False)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return [], []
    headers = [_text(value) for value in rows[0]]
    records = []
    for values in rows[1:]:
        records.append({headers[i]: ("" if i >= len(values) or values[i] is None else values[i])
                        for i in range(len(headers)) if headers[i]})
    return headers, records


def load_env_file(path):
    values = {}
    if not path or not os.path.exists(path):
        return values
    with open(path, encoding="utf-8") as handle:
        for line in handle:
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            key, value = line.split("=", 1)
            values[key.strip()] = value.strip().strip('"').strip("'")
    return values


def load_config(path=None):
    """读取钉钉配置；环境变量优先，绝不在代码中内置真实凭证。"""
    config = {}
    if path:
        config.update(load_env_file(path))
    else:
        config_path = os.environ.get("DINGTALK_CONFIG")
        if not config_path:
            config_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "config.env")
        config.update(load_env_file(config_path))
    for key in ("DINGTALK_ROUTE_FILE", "DINGTALK_WEBHOOK", "DINGTALK_SECRET"):
        if os.environ.get(key):
            config[key] = os.environ[key]
    return config


def load_routes(path):
    if not path or not os.path.exists(path):
        return {}
    with open(path, encoding="utf-8") as handle:
        value = json.load(handle)
    if not isinstance(value, dict):
        raise ValueError("钉钉路由文件必须是 JSON 对象")
    # 推荐格式为 {"routes": {"通知人": {"webhook": "https://..."}}}；
    # 同时兼容直接以通知人作为键的轻量格式。
    routes = value.get("routes") or value.get("通知人") or value
    return routes if isinstance(routes, dict) else {}


def _route_value(routes, notifier):
    normalized = re.sub(r"\s+", "", _text(notifier))
    for key, value in routes.items():
        if key in ("default", "routes", "通知人"):
            continue
        if re.sub(r"\s+", "", _text(key)) == normalized:
            return value
    # 不允许 default 兜底。通知人没有一对一路由时必须跳过，避免不同企业
    # 的政策结果被汇总广播到同一个群。
    return None


def resolve_route(routes, notifier, global_config=None):
    raw = _route_value(routes, notifier)
    if raw is None:
        return None
    if isinstance(raw, str):
        route = {"webhook": raw}
    elif isinstance(raw, dict):
        route = dict(raw)
    else:
        return None
    if route.get("use_config") or route.get("webhook") == "global":
        route["webhook"] = (global_config or {}).get("DINGTALK_WEBHOOK", "")
    if not route.get("secret") and global_config:
        route["secret"] = global_config.get("DINGTALK_SECRET", "")
    webhook = _text(route.get("webhook"))
    parsed = urlparse(webhook)
    host = (parsed.hostname or "").lower()
    if (not webhook or parsed.scheme.lower() != "https"
            or not (host == "dingtalk.com" or host.endswith(".dingtalk.com"))):
        return None
    route["webhook"] = webhook
    return route


def signed_webhook(webhook, secret, timestamp=None):
    """按钉钉自定义机器人签名规则生成带 timestamp/sign 的 URL。"""
    if not secret:
        return webhook
    timestamp = int(timestamp if timestamp is not None else time.time() * 1000)
    string_to_sign = f"{timestamp}\n{secret}"
    digest = hmac.new(secret.encode("utf-8"), string_to_sign.encode("utf-8"), hashlib.sha256).digest()
    sign = base64.b64encode(digest).decode("ascii")
    parsed = urlparse(webhook)
    params = parse_qsl(parsed.query, keep_blank_values=True)
    params.extend([("timestamp", str(timestamp)), ("sign", sign)])
    return urlunparse(parsed._replace(query=urlencode(params)))


def _candidate_rows(policies, enterprises, results, today):
    new_ids = {norm_key(row.get("序号", "")) for row in policies
               if _text(row.get("政策状态", "")) not in {"已过期", "已失效", "无效", "废止"}
               and (not _text(row.get("申报截止", ""))
                    or _text(row.get("申报截止", "")) >= today)}
    enterprise_by_name = {_text(row.get("企业名称", "")): row for row in enterprises}
    policy_by_id = {norm_key(row.get("序号", "")): row for row in policies}
    rows = []
    skipped = []
    for row in results:
        conclusion = _text(row.get("匹配结论", ""))
        if norm_key(row.get("政策ID", "")) not in new_ids:
            continue
        if conclusion not in NOTIFYABLE_CONCLUSIONS or _text(row.get("可通知", "是")) == "否":
            continue
        if _text(row.get("通知状态", "")) in SUCCESS_STATES:
            continue
        enterprise = enterprise_by_name.get(_text(row.get("企业名称", "")))
        notifier = _text((enterprise or {}).get("通知人", ""))
        if not notifier:
            skipped.append((result_key(row), "企业档案缺少通知人"))
            continue
        enriched = dict(row)
        enriched["通知人"] = notifier
        enriched["政策标题"] = _text((policy_by_id.get(norm_key(row.get("政策ID", ""))) or {}).get("标题", ""))
        enriched["原文链接"] = _text((policy_by_id.get(norm_key(row.get("政策ID", ""))) or {}).get("原文链接", ""))
        enriched["支持额度"] = _text((policy_by_id.get(norm_key(row.get("政策ID", ""))) or {}).get("支持额度", ""))
        enriched["申报截止"] = _text((policy_by_id.get(norm_key(row.get("政策ID", ""))) or {}).get("申报截止", ""))
        rows.append(enriched)
    return rows, skipped


def _message_for_rows(rows):
    lines = ["湖北/武汉企业补贴 · 待处理匹配结果", ""]
    for row in rows:
        note = _text(row.get("命中条件") if row.get("匹配结论") == "匹配" else row.get("缺口", ""))
        if len(note) > 180:
            note = note[:180] + "…"
        lines.append(f"【{row.get('企业名称', '')}】{row.get('匹配结论', '')}｜{row.get('政策标题', '')}")
        if note:
            lines.append(f"{note}")
        if row.get("支持额度"):
            lines.append(f"补贴金额：{row['支持额度']}")
        if row.get("申报截止"):
            lines.append(f"申报截止：{row['申报截止']}")
        if row.get("原文链接"):
            lines.append(f"原文：{row['原文链接']}")
        lines.append(f"通知编号：{result_key(row)}")
        lines.append("")
    return "\n".join(lines).strip()


def send_group(route, rows, dry_run=True, request_post=None):
    content = _message_for_rows(rows)
    if dry_run:
        print(f"[dry-run] 群路由 {route.get('name') or '已配置群'}: {len(rows)} 条")
        print(content)
        return True
    webhook = signed_webhook(route["webhook"], _text(route.get("secret")))
    post = request_post or requests.post
    response = post(webhook, json={"msgtype": "text", "text": {"content": content}}, timeout=TIMEOUT)
    if not (200 <= response.status_code < 300):
        raise RuntimeError(f"钉钉通知失败: HTTP {response.status_code}")
    try:
        body = response.json()
    except ValueError:
        body = {}
    if not isinstance(body, dict) or body.get("errcode") != 0:
        raise RuntimeError("钉钉通知未返回明确业务成功回执，请核对发送结果")
    return True


def write_notification_status(path, keys, status="已推送"):
    if not keys or not os.path.exists(path):
        return 0
    workbook = openpyxl.load_workbook(path)
    sheet = workbook.active
    rows = list(sheet.iter_rows())
    if not rows:
        return 0
    headers = [_text(cell.value) for cell in rows[0]]
    try:
        policy_col = headers.index("政策ID")
        enterprise_col = headers.index("企业名称")
    except ValueError as exc:
        raise RuntimeError("匹配结果表缺少政策ID或企业名称列") from exc
    if "通知状态" not in headers:
        status_col = len(headers)
        sheet.cell(row=1, column=status_col + 1, value="通知状态")
    else:
        status_col = headers.index("通知状态")
    changed = 0
    for row_number, row in enumerate(rows[1:], start=2):
        key = f"{norm_key(row[policy_col].value)}||{norm_key(row[enterprise_col].value)}"
        cell = row[status_col] if status_col < len(row) else sheet.cell(row=row_number, column=status_col + 1)
        if key in keys and cell.value != status:
            cell.value = status
            changed += 1
    if changed:
        workbook.save(path)
    return changed


def main(argv=None):
    parser = argparse.ArgumentParser(description="钉钉群级通知适配器（默认 dry-run）")
    parser.add_argument("--results", default="match_results.xlsx")
    parser.add_argument("--policies", default="subsidy_records.xlsx")
    parser.add_argument("--enterprises", default="enterprises.xlsx")
    parser.add_argument("--routes", default="")
    parser.add_argument("--config", default="")
    parser.add_argument("--today", default="", help="测试用日期 YYYY-MM-DD，默认今天")
    parser.add_argument("--dry-run", action="store_true", help="仅打印，不发送（默认行为）")
    parser.add_argument("--send", action="store_true", help="显式允许真实群 Webhook 外发")
    args = parser.parse_args(argv)
    data_dir = os.environ.get("SUBSIDY_DATA_DIR") or os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "data")
    def resolve(path, default):
        path = path or default
        return path if os.path.isabs(path) else os.path.join(data_dir, path)
    results_path = resolve(args.results, "match_results.xlsx")
    policies_path = resolve(args.policies, "subsidy_records.xlsx")
    enterprises_path = resolve(args.enterprises, "enterprises.xlsx")
    config = load_config(args.config or None)
    route_path = args.routes or config.get("DINGTALK_ROUTE_FILE", "dingtalk_routes.json")
    route_path = resolve(route_path, "dingtalk_routes.json")
    routes = load_routes(route_path)
    if not routes:
        print("[跳过] 未配置钉钉路由文件，无安全的群级发送目标")
        return 0
    _, policies = read_xlsx(policies_path)
    _, enterprises = read_xlsx(enterprises_path)
    _, results = read_xlsx(results_path)
    today = args.today or datetime.date.today().strftime("%Y-%m-%d")
    candidates, skipped = _candidate_rows(policies, enterprises, results, today)
    ledger = Ledger(os.path.dirname(results_path))
    blocked = ledger.blocked()
    candidates = [row for row in candidates if delivery_key(row["政策ID"], row["企业名称"]) not in blocked]
    for key, reason in skipped:
        print(f"[跳过] {key}: {reason}")
    if not candidates:
        print("[跳过] 当前无可自动发送的匹配结果；发送结果不明项请查看通知台账")
        return 0
    grouped = {}
    for row in candidates:
        route = resolve_route(routes, row["通知人"], config)
        if not route:
            print(f"[跳过] {result_key(row)}: 通知人无有效 HTTPS 群路由")
            continue
        route_id = (route["webhook"], _text(route.get("secret")))
        grouped.setdefault(route_id, {"route": route, "rows": []})["rows"].append(row)
    if not grouped:
        print("[跳过] 当前结果没有可用钉钉群路由")
        return 0
    # dry-run 是默认值；环境变量不能绕过显式 --send 保护。
    dry_run = args.dry_run or not args.send
    sent_keys = set()
    for group in grouped.values():
        keys = {delivery_key(row["政策ID"], row["企业名称"]) for row in group["rows"]}
        if not dry_run:
            ledger.reserve(keys)
        send_group(group["route"], group["rows"], dry_run=dry_run)
        if not dry_run:
            ledger.finish(keys, "dingtalk webhook errcode=0")
            sent_keys.update(result_key(row) for row in group["rows"])
            # Persist each successful group before moving to the next one.
            write_notification_status(results_path, {result_key(row) for row in group["rows"]})
    if sent_keys:
        write_notification_status(results_path, sent_keys)
        print(f"[ok] 成功发送 {len(sent_keys)} 条，已回写通知状态")
    else:
        print(f"[dry-run] 共 {sum(len(v['rows']) for v in grouped.values())} 条，未回写通知状态")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
飞书多维表格同步 + 通知 (Skill 阶段6)
=================================================================
数据中心在飞书多维表格，本地 xlsx 是缓存/备份。任何智能体拿到
"凭证 + 多维表格位置"即可驱动本脚本，不绑定特定平台。

用法:
  python feishu_sync.py --check      # 自检: 凭证/表格定位/字段清单(接入第一步)
  python feishu_sync.py --sync       # 本地 → 飞书: 政策表/匹配表 upsert
  python feishu_sync.py --pull       # 飞书 → 本地: 企业档案(飞书为维护入口)
  python feishu_sync.py --notify     # 推送今日摘要 + 匹配结果到飞书群
  python feishu_sync.py --pull --sync --notify   # 每日自动化全套（先拉企业档案）

配置(优先级: 环境变量 > SUBSIDY_CONFIG_FILE 指定文件或项目根 config.env):
  FEISHU_APP_ID / FEISHU_APP_SECRET   自建应用凭证(授予多维表格读写权限)
  FEISHU_BASE_URL                     多维表格 URL 或 app_token(只给这个即可,
                                      三张表按表名自动定位: 政策清单/企业档案/匹配结果)
  FEISHU_POLICY_TABLE 等              (可选) 显式指定 table_id 覆盖自动定位
  FEISHU_WEBHOOK                      (可选) 群机器人 Webhook, 用于推送

同步规则:
  - 政策表: 按「原文链接」upsert; 人工列(申报截止/处理状态/是否匹配/备注)
    仅当飞书侧为空时才写入, 不覆盖人工已填内容
  - 匹配表: 按「政策ID + 企业名称」复合键 upsert
  - 企业档案: 以飞书为准, --pull 拉回本地供匹配引擎使用
依赖: requests openpyxl
"""
import os, re, sys, json, argparse, datetime
import openpyxl
import requests
from notification_ledger import Ledger, key as delivery_key
from setup_state import read_config

BASE = "https://open.feishu.cn/open-apis"
TIMEOUT = 25

CONFIG_KEYS = ["FEISHU_APP_ID", "FEISHU_APP_SECRET", "FEISHU_BASE_URL",
               "FEISHU_APP_TOKEN", "FEISHU_POLICY_TABLE", "FEISHU_ENT_TABLE",
               "FEISHU_MATCH_TABLE", "FEISHU_WEBHOOK", "SKILL_MODE"]

# 自动定位: 表名包含关键词即认为是对应数据表
TABLE_NAME_HINTS = {
    "FEISHU_POLICY_TABLE": ["政策", "清单"],
    "FEISHU_ENT_TABLE": ["企业", "档案"],
    "FEISHU_MATCH_TABLE": ["匹配"],
}

POLICY_KEY = "原文链接"
POLICY_HUMAN_COLS = {"申报截止", "处理状态", "是否匹配", "备注"}  # 人工列: 只补空不覆盖


# ===========================================================================
# 配置与定位
# ===========================================================================
def load_config():
    """Use the same explicit configuration path as setup and daily execution."""
    return read_config()


def app_token_of(cfg):
    """从 FEISHU_BASE_URL(完整URL) 或 FEISHU_APP_TOKEN 提取 app_token。"""
    raw = cfg.get("FEISHU_BASE_URL") or cfg.get("FEISHU_APP_TOKEN") or ""
    m = re.search(r'base/([A-Za-z0-9]+)', raw)
    return m.group(1) if m else raw.strip()


def get_token(cfg):
    r = requests.post(f"{BASE}/auth/v3/tenant_access_token/internal",
                      json={"app_id": cfg["FEISHU_APP_ID"],
                            "app_secret": cfg["FEISHU_APP_SECRET"]},
                      timeout=TIMEOUT)
    d = r.json()
    if d.get("code") != 0:
        raise RuntimeError(f"获取 token 失败(检查 app_id/secret): {d}")
    return d["tenant_access_token"]


def _hdr(token):
    return {"Authorization": f"Bearer {token}",
            "Content-Type": "application/json"}


def _get(token, path, params=None):
    r = requests.get(f"{BASE}{path}", headers=_hdr(token), params=params,
                     timeout=TIMEOUT)
    d = r.json()
    if d.get("code") != 0:
        raise RuntimeError(f"GET {path} 失败: {d.get('msg')} (code={d.get('code')})")
    return d.get("data", {})


def list_tables(token, app_token):
    data = _get(token, f"/bitable/v1/apps/{app_token}/tables",
                params={"page_size": 100})
    return data.get("items", [])


def resolve_tables(token, cfg, app_token):
    """显式 table_id 优先, 否则按表名关键词自动定位。"""
    resolved = {}
    tables = list_tables(token, app_token)
    for key, hints in TABLE_NAME_HINTS.items():
        if cfg.get(key):
            resolved[key] = cfg[key]
            continue
        for t in tables:
            name = t.get("name", "")
            if all(h in name for h in hints):
                resolved[key] = t["table_id"]
                break
    return resolved, tables


def list_fields(token, app_token, table_id):
    """返回 {字段名: ui_type}，如 Text/Number/SingleSelect/MultiSelect/DateTime/Url。"""
    data = _get(token, f"/bitable/v1/apps/{app_token}/tables/{table_id}/fields",
                params={"page_size": 100})
    return {f["field_name"]: f.get("ui_type", "Text") for f in data.get("items", [])}


# ===========================================================================
# 值转换: 本地 <-> 飞书字段类型
# ===========================================================================
def to_feishu(v, ui_type):
    """本地 xlsx 值 → 飞书字段值。None 表示不写入该字段。"""
    if v is None:
        return None
    if hasattr(v, "strftime"):
        v = v.strftime("%Y-%m-%d %H:%M" if (v.hour or v.minute) else "%Y-%m-%d")
    s = str(v).strip()
    if not s:
        return None
    if ui_type == "Number":
        try:
            return float(s)
        except ValueError:
            return None
    if ui_type == "DateTime":
        for fmt in ("%Y-%m-%d %H:%M", "%Y-%m-%d"):
            try:
                dt = datetime.datetime.strptime(s, fmt)
                return int(dt.timestamp() * 1000)
            except ValueError:
                continue
        return None
    if ui_type == "MultiSelect":
        return [x for x in re.split(r'[、，,;/；]', s) if x]
    if ui_type == "Url":
        return {"link": s, "text": s}
    return s


def _flatten_text(v):
    """飞书文本字段读出的是 [{text:...}] 段结构，拼成字符串。"""
    if isinstance(v, list):
        return "".join(seg.get("text", "") if isinstance(seg, dict) else str(seg)
                       for seg in v)
    return v


def from_feishu(v, ui_type):
    """飞书字段值 → 本地纯值。"""
    if v is None:
        return ""
    if ui_type == "DateTime":
        try:
            return datetime.datetime.fromtimestamp(int(v) / 1000).strftime("%Y-%m-%d")
        except (TypeError, ValueError, OSError):
            return str(v)
    if ui_type == "MultiSelect":
        return "、".join(str(x) for x in v) if isinstance(v, list) else str(v)
    if ui_type == "Url":
        return v.get("link", "") if isinstance(v, dict) else str(v)
    if ui_type == "Number":
        if isinstance(v, float) and v.is_integer():
            return int(v)
        return v
    return _flatten_text(v)


def norm_key(v):
    """主键归一化: 12.0 -> '12'，去首尾空格。"""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _result_key(row):
    return f"{norm_key(row.get('政策ID', ''))}||{norm_key(row.get('企业名称', ''))}"


# ===========================================================================
# 记录读取(全量一次, 避免每行一次 search 调用)
# ===========================================================================
def fetch_all_records(token, app_token, table_id):
    items, page_token = [], None
    while True:
        params = {"page_size": 500}
        if page_token:
            params["page_token"] = page_token
        data = _get(token, f"/bitable/v1/apps/{app_token}/tables/{table_id}/records",
                    params=params)
        items.extend(data.get("items", []))
        if not data.get("has_more"):
            break
        page_token = data.get("page_token")
    return items


def read_xlsx(path):
    if not os.path.exists(path):
        return [], []
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []
    hdr = [str(h) for h in rows[0]]
    data = [{hdr[i]: ("" if (i >= len(r) or r[i] is None) else r[i])
             for i in range(len(hdr))} for r in rows[1:]]
    return hdr, data


# ===========================================================================
# 同步: 本地 → 飞书
# ===========================================================================
def upsert_table(token, app_token, table_id, xlsx_path, key_fn, human_cols=frozenset()):
    """key_fn(row_dict)->str 生成去重键; human_cols 仅当飞书侧为空时才写入。"""
    hdr, records = read_xlsx(xlsx_path)
    if not records:
        print(f"  [跳过] 本地表为空或不存在: {xlsx_path}")
        return
    meta = list_fields(token, app_token, table_id)
    existing = fetch_all_records(token, app_token, table_id)
    idx = {}
    for rec in existing:
        flat = {k: from_feishu(v, meta.get(k, "Text"))
                for k, v in rec.get("fields", {}).items()}
        idx[key_fn(flat)] = (rec["record_id"], flat)

    cre, upd = [], []
    skipped_human = 0
    for row in records:
        kv = key_fn(row)
        if not kv:
            continue
        fields = {}
        for k, v in row.items():
            if k not in meta:      # 飞书表里没有的列直接跳过
                continue
            fv = to_feishu(v, meta[k])
            if fv is not None:
                fields[k] = fv
        if kv in idx:
            record_id, flat = idx[kv]
            for col in human_cols:
                if col in fields and str(flat.get(col, "")).strip():
                    del fields[col]   # 飞书侧已有人工内容, 不覆盖
                    skipped_human += 1
            if fields:
                upd.append({"record_id": record_id, "fields": fields})
        else:
            cre.append({"fields": fields})

    ok = 0
    failures = []
    for batch in (cre[i:i + 100] for i in range(0, len(cre), 100)):
        d = requests.post(f"{BASE}/bitable/v1/apps/{app_token}/tables/{table_id}"
                          f"/records/batch_create", headers=_hdr(token),
                          json={"records": batch}, timeout=TIMEOUT).json()
        if d.get("code") == 0:
            ok += len(batch)
        else:
            failures.append(f"batch_create: {d.get('msg')}")
    for batch in (upd[i:i + 100] for i in range(0, len(upd), 100)):
        d = requests.post(f"{BASE}/bitable/v1/apps/{app_token}/tables/{table_id}"
                          f"/records/batch_update", headers=_hdr(token),
                          json={"records": batch}, timeout=TIMEOUT).json()
        if d.get("code") == 0:
            ok += len(batch)
        else:
            failures.append(f"batch_update: {d.get('msg')}")
    print(f"  同步 {os.path.basename(xlsx_path)}: 新建 {len(cre)} / 更新 {len(upd)}"
          + (f" / 保留人工列 {skipped_human} 处" if skipped_human else ""))
    if failures:
        raise RuntimeError("飞书同步失败: " + "；".join(failures[:5]))
    return ok


def sync_all(token, cfg, app_token, tables, only, data_dir):
    if only in ("all", "policy"):
        if not tables.get("FEISHU_POLICY_TABLE"):
            raise RuntimeError("未定位到飞书政策清单表，拒绝假成功")
        upsert_table(token, app_token, tables["FEISHU_POLICY_TABLE"],
                     os.path.join(data_dir, "subsidy_records.xlsx"),
                     key_fn=lambda r: norm_key(r.get(POLICY_KEY, "")),
                     human_cols=POLICY_HUMAN_COLS)
    if only in ("all", "match"):
        if not tables.get("FEISHU_MATCH_TABLE"):
            raise RuntimeError("未定位到飞书匹配结果表，拒绝假成功")
        upsert_table(token, app_token, tables["FEISHU_MATCH_TABLE"],
                     os.path.join(data_dir, "match_results.xlsx"),
                     key_fn=lambda r: f"{norm_key(r.get('政策ID',''))}||{norm_key(r.get('企业名称',''))}")


# ===========================================================================
# 拉回: 飞书 → 本地(企业档案以飞书为准)
# ===========================================================================
def pull_table(token, app_token, table_id, xlsx_path):
    meta = list_fields(token, app_token, table_id)
    records = fetch_all_records(token, app_token, table_id)
    wb = openpyxl.Workbook()
    ws = wb.active
    cols = list(meta.keys())
    ws.append(cols)
    for rec in records:
        flat = {k: from_feishu(v, meta.get(k, "Text"))
                for k, v in rec.get("fields", {}).items()}
        ws.append([flat.get(c, "") for c in cols])
    wb.save(xlsx_path)
    print(f"  拉回 {len(records)} 行 → {os.path.basename(xlsx_path)}")


# ===========================================================================
# 通知
# ===========================================================================
NOTIFYABLE_CONCLUSIONS = {"匹配"}
SUCCESS_NOTIFICATION_STATES = {"已推送", "已发送"}


def _write_notification_status(data_dir, keys, status="已推送"):
    """只更新本地匹配表的通知列，保留旧表的其它列和人工处理状态。"""
    path = os.path.join(data_dir, "match_results.xlsx")
    if not os.path.exists(path) or not keys:
        return 0
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    rows = list(ws.iter_rows())
    if not rows:
        return 0
    headers = [str(cell.value or "") for cell in rows[0]]
    try:
        key_col = headers.index("政策ID")
        ent_col = headers.index("企业名称")
    except ValueError:
        raise RuntimeError("匹配结果表缺少政策ID或企业名称列")
    if "通知状态" in headers:
        status_col = headers.index("通知状态")
    else:
        status_col = len(headers)
        ws.cell(row=1, column=status_col + 1, value="通知状态")
        headers.append("通知状态")
    changed = 0
    for row_number, row in enumerate(rows[1:], start=2):
        key = f"{norm_key(row[key_col].value)}||{norm_key(row[ent_col].value)}"
        if key in keys:
            cell = row[status_col] if status_col < len(row) else ws.cell(row=row_number, column=status_col + 1)
            if cell.value != status:
                cell.value = status
                changed += 1
    if changed:
        wb.save(path)
    return changed


def notify(token, cfg, data_dir, limit=50, dry_run=False):
    """发送当前有效且尚未推送的匹配结果，支持隔天补档和失败补偿。"""
    webhook = cfg.get("FEISHU_WEBHOOK")
    if not webhook:
        print("  [跳过] 未配置 FEISHU_WEBHOOK")
        return 0
    if not 1 <= limit <= 50:
        raise ValueError("通知批次必须为 1 至 50 条")
    _, policies = read_xlsx(os.path.join(data_dir, "subsidy_records.xlsx"))
    _, matches = read_xlsx(os.path.join(data_dir, "match_results.xlsx"))
    today = datetime.date.today().strftime("%Y-%m-%d")
    new_ids = {norm_key(p.get("序号", "")) for p in policies
               if str(p.get("政策状态", "")).strip() not in {"已过期", "已失效", "无效", "废止"}
               and (not str(p.get("申报截止", "")).strip()
                    or str(p.get("申报截止", "")).strip() >= today)}
    policy_by_id = {norm_key(p.get("序号", "")): p for p in policies}
    ledger = Ledger(data_dir)
    blocked = ledger.blocked()
    candidates = []
    for row in matches:
        conclusion = str(row.get("匹配结论", "")).strip()
        allowed = str(row.get("可通知", "是")).strip() != "否"
        if (norm_key(row.get("政策ID", "")) in new_ids
                and conclusion in NOTIFYABLE_CONCLUSIONS
                and allowed
                and str(row.get("通知状态", "")).strip() not in SUCCESS_NOTIFICATION_STATES
                and delivery_key(row.get("政策ID"), row.get("企业名称")) not in blocked):
            candidates.append(row)
    if not candidates:
        print("  [跳过] 当前无可自动发送的匹配结果；发送结果不明项请查看通知台账")
        return 0

    # Only acknowledge rows actually included in this card; remaining rows stay pending.
    total = len(candidates)
    candidates = candidates[:limit]
    lines = ["湖北/武汉企业补贴 · 待处理匹配结果", ""]
    for row in candidates:
        pid = norm_key(row.get("政策ID", ""))
        policy = policy_by_id.get(pid, {})
        note = str((row.get("命中条件") if row.get("匹配结论") == "匹配" else row.get("缺口")) or "")
        if len(note) > 160:
            note = note[:160] + "…"
        lines.append(f"- {row.get('企业名称', '')}：{row.get('匹配结论', '')}｜{policy.get('标题', '')}")
        if note:
            lines.append(f"  {note}")
        if str(policy.get("支持额度", "")).strip():
            lines.append(f"  补贴金额：{policy.get('支持额度')}")
        if str(policy.get("申报截止", "")).strip():
            lines.append(f"  申报截止：{policy.get('申报截止')}")
        if policy.get("原文链接"):
            lines.append(f"  原文：{policy['原文链接']}")
    card = {"msg_type": "interactive",
            "card": {"header": {"title": {"tag": "plain_text",
                                          "content": "湖北/武汉企业补贴 · 每日播报"}},
                     "elements": [{"tag": "div",
                                   "text": {"tag": "lark_md",
                                             "content": "\n".join(lines)}}]}}
    if dry_run:
        print(json.dumps({"pending_count": total, "batch_count": len(candidates),
                          "match_keys": [delivery_key(row.get("政策ID"), row.get("企业名称")) for row in candidates],
                          "preview": card}, ensure_ascii=False, indent=2))
        return len(candidates)
    delivery_keys = {delivery_key(row.get("政策ID"), row.get("企业名称")) for row in candidates}
    ledger.reserve(delivery_keys)
    response = requests.post(webhook, json=card, timeout=TIMEOUT)
    if not (200 <= response.status_code < 300):
        raise RuntimeError(f"飞书通知失败: HTTP {response.status_code}")
    try:
        body = response.json()
    except ValueError:
        body = {}
    if not isinstance(body, dict) or body.get("code") != 0:
        raise RuntimeError("飞书通知未返回明确业务成功回执，请核对发送结果")
    ledger.finish(delivery_keys, "feishu webhook code=0")
    sent_keys = {_result_key(row) for row in candidates}
    _write_notification_status(data_dir, sent_keys)
    print(f"  推送卡片成功: {len(candidates)} 条，已回写通知状态")
    return len(candidates)


# ===========================================================================
# 自检
# ===========================================================================
def check(cfg):
    if not cfg.get("FEISHU_APP_ID") or not cfg.get("FEISHU_APP_SECRET"):
        print("[!] 未配置凭证。接入步骤：")
        print("  1. open.feishu.cn 创建企业自建应用，拿 app_id / app_secret")
        print("  2. 应用开通「多维表格」读写权限，并把应用加进目标表(表内 ... → 添加应用)")
        print("  3. 复制 config.example.env 为 config.env，填入凭证 + 多维表格 URL")
        print("  4. 再运行本命令验证")
        return False
    token = get_token(cfg)
    print("[ok] tenant_access_token 获取成功")
    app_token = app_token_of(cfg)
    if not app_token:
        print("[!] 未配置 FEISHU_BASE_URL（多维表格 URL 或 app_token）")
        return False
    tables, all_tables = resolve_tables(token, cfg, app_token)
    print(f"[ok] 多维表格可访问，共 {len(all_tables)} 张数据表：")
    for t in all_tables:
        print(f"     - {t.get('name')} ({t['table_id']})")
    ok = True
    for key, hints in TABLE_NAME_HINTS.items():
        tid = tables.get(key)
        label = {"FEISHU_POLICY_TABLE": "政策清单", "FEISHU_ENT_TABLE": "企业档案",
                 "FEISHU_MATCH_TABLE": "匹配结果"}[key]
        if not tid:
            print(f"[!] 未定位到「{label}」表（表名需含 {'+'.join(hints)}，"
                  f"或在 config.env 显式配置 {key}）")
            ok = False
            continue
        meta = list_fields(token, app_token, tid)
        print(f"[ok] {label}表 {tid}，字段 {len(meta)} 个：")
        for name, ui in meta.items():
            print(f"       {name} ({ui})")
    print(f"[{'ok' if cfg.get('FEISHU_WEBHOOK') else '--'}] 群机器人 Webhook "
          f"{'已配置' if cfg.get('FEISHU_WEBHOOK') else '未配置(可选)'}")
    return ok


# ===========================================================================
def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--check", action="store_true")
    ap.add_argument("--sync", action="store_true")
    ap.add_argument("--pull", action="store_true")
    ap.add_argument("--notify", action="store_true")
    ap.add_argument("--notify-preview", action="store_true", help="仅预览通知，不访问平台或预留发送")
    ap.add_argument("--notify-limit", type=int, default=50, help="本轮通知条数，首次测试可设为 1")
    ap.add_argument("--only", default="all", choices=["all", "policy", "match"])
    args = ap.parse_args(argv)
    data_dir = os.environ.get("SUBSIDY_DATA_DIR") or os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "data")
    cfg = load_config()

    if args.notify_preview:
        notify("", cfg, data_dir, args.notify_limit, True)
        return 0

    if args.check or not (args.sync or args.pull or args.notify):
        return 0 if check(cfg) else 2
    if cfg.get("SKILL_MODE", "").lower() in ("local", "dingtalk"):
        print("[跳过] 当前模式不使用飞书同步")
        return 0
    if not cfg.get("FEISHU_APP_ID") or not cfg.get("FEISHU_APP_SECRET") or not app_token_of(cfg):
        raise RuntimeError("未配置完整飞书凭证/多维表格位置；请先运行 --check")
    token = get_token(cfg)
    app_token = app_token_of(cfg)
    tables, _ = resolve_tables(token, cfg, app_token)
    if args.pull:
        if not tables.get("FEISHU_ENT_TABLE"):
            raise RuntimeError("未定位到飞书企业档案表，拒绝使用旧企业档案继续运行")
        pull_table(token, app_token, tables["FEISHU_ENT_TABLE"],
                   os.path.join(data_dir, "enterprises.xlsx"))
    # 调用方传入 --pull --sync --notify 时严格保持 pull → sync → notify 顺序。
    if args.sync:
        sync_all(token, cfg, app_token, tables, args.only, data_dir)
    if args.notify:
        notify(token, cfg, data_dir, args.notify_limit)
        # 通知成功后本地状态已变为“已推送”，再同步一次匹配表使飞书状态一致。
        if args.sync and args.only in ("all", "match"):
            sync_all(token, cfg, app_token, tables, "match", data_dir)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

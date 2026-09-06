#!/usr/bin/env python3
"""提取并校验 DOCX 的结构保真基线。

V1.5 将参考 DOCX 视为母文件。该脚本记录正文元素、章节、表格、分节、
页眉页脚及媒体关系，并可在交付前把成稿与参考清单逐项比对。
"""

from __future__ import annotations

import argparse
import json
import re
import sys
import zipfile
from collections import Counter
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


HEADING_NAMES = {f"Heading {level}": level for level in range(1, 10)}
HEADING_NAMES.update({f"标题 {level}": level for level in range(1, 10)})
CHAPTER_TEXT = re.compile(r"^第\s*[一二三四五六七八九十百零〇0-9]+\s*章")


def _clean(text: str) -> str:
    return re.sub(r"\s+", " ", text or "").strip()


def _paragraph_level(paragraph: Paragraph) -> int | None:
    style_name = paragraph.style.name if paragraph.style is not None else ""
    if style_name in HEADING_NAMES:
        return HEADING_NAMES[style_name]
    outline = paragraph._p.xpath("./w:pPr/w:outlineLvl/@w:val")
    if outline:
        try:
            return int(outline[0]) + 1
        except ValueError:
            pass
    if CHAPTER_TEXT.match(_clean(paragraph.text)):
        return 1
    return None


def _body_elements(doc: Document) -> list[dict[str, Any]]:
    elements: list[dict[str, Any]] = []
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            paragraph = Paragraph(child, doc)
            text = _clean(paragraph.text)
            elements.append(
                {
                    "kind": "paragraph",
                    "text": text,
                    "style": paragraph.style.name if paragraph.style is not None else "",
                    "heading_level": _paragraph_level(paragraph),
                }
            )
        elif child.tag == qn("w:tbl"):
            table = Table(child, doc)
            elements.append(
                {
                    "kind": "table",
                    "rows": len(table.rows),
                    "cols": max((len(row.cells) for row in table.rows), default=0),
                    "header": [_clean(cell.text) for cell in table.rows[0].cells] if table.rows else [],
                }
            )
    return elements


def _chapter_metrics(elements: list[dict[str, Any]]) -> list[dict[str, Any]]:
    starts = [
        index
        for index, element in enumerate(elements)
        if element["kind"] == "paragraph" and element.get("heading_level") == 1
    ]
    chapters: list[dict[str, Any]] = []
    for position, start in enumerate(starts):
        end = starts[position + 1] if position + 1 < len(starts) else len(elements)
        chunk = elements[start:end]
        paragraphs = [element for element in chunk if element["kind"] == "paragraph"]
        tables = [element for element in chunk if element["kind"] == "table"]
        chapters.append(
            {
                "title": chunk[0]["text"],
                "paragraphs": len(paragraphs),
                "nonempty_paragraphs": sum(bool(element["text"]) for element in paragraphs),
                "tables": len(tables),
                "headings": [
                    {"level": element["heading_level"], "text": element["text"]}
                    for element in paragraphs
                    if element.get("heading_level")
                ],
            }
        )
    return chapters


def _story_signature(story) -> dict[str, Any]:
    return {
        "paragraphs": len(story.paragraphs),
        "tables": len(story.tables),
        "text": _clean("\n".join(paragraph.text for paragraph in story.paragraphs)),
    }


def build_manifest(path: str | Path) -> dict[str, Any]:
    source = Path(path).resolve()
    doc = Document(source)
    elements = _body_elements(doc)
    paragraph_elements = [element for element in elements if element["kind"] == "paragraph"]
    table_elements = [element for element in elements if element["kind"] == "table"]
    with zipfile.ZipFile(source) as package:
        package_parts = package.namelist()
        media_parts = sorted(name for name in package_parts if name.startswith("word/media/"))
        headers = sorted(name for name in package_parts if re.fullmatch(r"word/header\d+\.xml", name))
        footers = sorted(name for name in package_parts if re.fullmatch(r"word/footer\d+\.xml", name))
        comments = sorted(name for name in package_parts if name.startswith("word/comments"))

    manifest: dict[str, Any] = {
        "schema_version": 1,
        "source": str(source),
        "body": {
            "elements": len(elements),
            "paragraphs": len(paragraph_elements),
            "nonempty_paragraphs": sum(bool(element["text"]) for element in paragraph_elements),
            "tables": len(table_elements),
            "sections": len(doc.sections),
            "inline_shapes": len(doc.inline_shapes),
            "textboxes": len(doc.element.xpath(".//w:txbxContent")),
            "drawings": len(doc.element.xpath(".//w:drawing")),
            "pict": len(doc.element.xpath(".//w:pict")),
        },
        "chapters": _chapter_metrics(elements),
        "tables": [
            {"index": index + 1, **table}
            for index, table in enumerate(table_elements)
        ],
        "sections": [
            {
                "index": index + 1,
                "start_type": str(section.start_type),
                "page_width_emu": int(section.page_width or 0),
                "page_height_emu": int(section.page_height or 0),
                "left_margin_emu": int(section.left_margin or 0),
                "right_margin_emu": int(section.right_margin or 0),
                "top_margin_emu": int(section.top_margin or 0),
                "bottom_margin_emu": int(section.bottom_margin or 0),
                "different_first_page": bool(section.different_first_page_header_footer),
                "header": _story_signature(section.header),
                "footer": _story_signature(section.footer),
                "first_page_header": _story_signature(section.first_page_header),
                "first_page_footer": _story_signature(section.first_page_footer),
            }
            for index, section in enumerate(doc.sections)
        ],
        "package": {
            "media_count": len(media_parts),
            "media_extensions": dict(sorted(Counter(Path(name).suffix.lower() for name in media_parts).items())),
            "header_parts": len(headers),
            "footer_parts": len(footers),
            "comment_parts": len(comments),
        },
    }
    return manifest


def compare_manifests(
    reference: dict[str, Any],
    candidate: dict[str, Any],
    *,
    allow_growth: bool = False,
) -> tuple[list[str], list[str]]:
    errors: list[str] = []
    warnings: list[str] = []

    for key in ("paragraphs", "tables", "sections"):
        expected = reference["body"][key]
        actual = candidate["body"][key]
        mismatch = actual < expected if allow_growth and key in ("paragraphs", "tables") else actual != expected
        if mismatch:
            comparator = "至少" if allow_growth and key in ("paragraphs", "tables") else "应为"
            errors.append(f"body.{key}: 参考{comparator} {expected}，成稿 {actual}")

    reference_chapters = reference["chapters"]
    candidate_chapters = candidate["chapters"]
    chapter_count_mismatch = (
        len(candidate_chapters) < len(reference_chapters)
        if allow_growth
        else len(candidate_chapters) != len(reference_chapters)
    )
    if chapter_count_mismatch:
        errors.append(f"章节数: 参考{'至少 ' if allow_growth else ''}{len(reference_chapters)}，成稿 {len(candidate_chapters)}")
    for index, expected in enumerate(reference_chapters):
        if index >= len(candidate_chapters):
            break
        actual = candidate_chapters[index]
        if _clean(actual["title"]) != _clean(expected["title"]):
            errors.append(f"第 {index + 1} 章标题: 参考“{expected['title']}”，成稿“{actual['title']}”")
        for key in ("paragraphs", "tables"):
            mismatch = actual[key] < expected[key] if allow_growth else actual[key] != expected[key]
            if mismatch:
                errors.append(
                    f"{expected['title']} {key}: 参考{'至少 ' if allow_growth else ''}{expected[key]}，成稿 {actual[key]}"
                )

    reference_tables = reference["tables"]
    candidate_tables = candidate["tables"]
    if not allow_growth and len(reference_tables) == len(candidate_tables):
        for expected, actual in zip(reference_tables, candidate_tables):
            if (expected["rows"], expected["cols"]) != (actual["rows"], actual["cols"]):
                errors.append(
                    f"表{expected['index']}尺寸: 参考 {expected['rows']}x{expected['cols']}，"
                    f"成稿 {actual['rows']}x{actual['cols']}"
                )

    for key in ("media_count", "header_parts", "footer_parts"):
        expected = reference["package"][key]
        actual = candidate["package"][key]
        if actual < expected:
            warnings.append(f"package.{key}: 参考 {expected}，成稿 {actual}")
    return errors, warnings


def _load_manifest(path: str | Path) -> dict[str, Any]:
    return json.loads(Path(path).read_text(encoding="utf-8-sig"))


def _write_manifest(manifest: dict[str, Any], path: str | Path) -> None:
    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(description="DOCX 结构清单与保真门禁")
    parser.add_argument("docx", help="参考或成稿 DOCX")
    parser.add_argument("--out", help="输出 manifest JSON")
    parser.add_argument("--compare", help="参考 manifest JSON 或参考 DOCX")
    parser.add_argument("--allow-growth", action="store_true", help="允许成稿增加段落或表格，但不允许缩水")
    args = parser.parse_args()

    manifest = build_manifest(args.docx)
    if args.out:
        _write_manifest(manifest, args.out)
        print(f"OK 结构清单: {Path(args.out).resolve()}")

    if args.compare:
        compare_path = Path(args.compare)
        reference = _load_manifest(compare_path) if compare_path.suffix.lower() == ".json" else build_manifest(compare_path)
        errors, warnings = compare_manifests(reference, manifest, allow_growth=args.allow_growth)
        print(f"结构门禁: 错误 {len(errors)} | 警告 {len(warnings)}")
        for error in errors:
            print(f"[错误] {error}")
        for warning in warnings:
            print(f"[警告] {warning}")
        if errors:
            sys.exit(1)
        print("PASS DOCX 主结构与参考一致")
    elif not args.out:
        print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()

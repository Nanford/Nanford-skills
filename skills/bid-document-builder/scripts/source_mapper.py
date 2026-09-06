#!/usr/bin/env python3
"""把项目 Excel 转成可追溯事实、单元格来源和算术冲突清单。"""

from __future__ import annotations

import argparse
import json
import re
import sys
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries


COLUMN_ALIASES = {
    "serial": ("序号", "编号"),
    "item": ("项目", "名称", "项目名称", "设备名称"),
    "feature": ("项目特征", "特征", "规格型号", "技术参数"),
    "unit": ("单位",),
    "quantity": ("数量", "工程量"),
    "unit_price": ("综合单价", "单价", "含税单价"),
    "total": ("合价", "金额", "总价", "小计"),
    "remark": ("备注",),
    "allocation": ("分摊", "公用部分", "承担部分"),
}


def _text(value: Any) -> str:
    return re.sub(r"\s+", "", str(value or "")).strip()


def _decimal(value: Any) -> Decimal | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value))
    cleaned = str(value).replace(",", "").replace("，", "").strip()
    match = re.fullmatch(r"[-+]?\d+(?:\.\d+)?", cleaned)
    if not match:
        return None
    try:
        return Decimal(cleaned)
    except InvalidOperation:
        return None


def _json_value(value: Any) -> Any:
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def _column_key(value: Any) -> str | None:
    normalized = _text(value)
    for key, aliases in COLUMN_ALIASES.items():
        if normalized in aliases:
            return key
    candidates: list[tuple[int, str]] = []
    for key, aliases in COLUMN_ALIASES.items():
        for alias in aliases:
            if alias in normalized:
                candidates.append((len(alias), key))
    return max(candidates, default=(0, ""))[1] or None


def _detect_header(sheet) -> tuple[int | None, dict[str, int]]:
    best_row: int | None = None
    best: dict[str, int] = {}
    for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 50)):
        mapping: dict[str, int] = {}
        for cell in row:
            key = _column_key(cell.value)
            if key and key not in mapping:
                mapping[key] = cell.column
        score = len(mapping)
        if score > len(best):
            best_row, best = row[0].row, mapping
    if len(best) < 3:
        return None, {}
    # 部分报价表把“分摊”等扩展列表头放在下一行，而主表头仍在首行。
    used_columns = set(best.values())
    for row in sheet.iter_rows(min_row=best_row + 1, max_row=min(sheet.max_row, best_row + 2)):
        for cell in row:
            key = _column_key(cell.value)
            if key and key not in best and cell.column not in used_columns:
                best[key] = cell.column
                used_columns.add(cell.column)
    return best_row, best


def _cell_value(raw_sheet, value_sheet, row: int, column: int) -> tuple[Any, str | None]:
    raw = raw_sheet.cell(row=row, column=column).value
    cached = value_sheet.cell(row=row, column=column).value
    formula = raw if isinstance(raw, str) and raw.startswith("=") else None
    return (cached if formula else raw), formula


def _evaluate_formula(formula: str, value_sheet) -> Decimal | None:
    expression = formula.strip().upper().replace("$", "")
    reference_match = re.fullmatch(r"=([A-Z]+\d+)", expression)
    if reference_match:
        return _decimal(value_sheet[reference_match.group(1)].value)
    sum_match = re.fullmatch(r"=SUM\(([A-Z]+\d+):([A-Z]+\d+)\)", expression)
    if sum_match:
        min_col, min_row, max_col, max_row = range_boundaries(f"{sum_match.group(1)}:{sum_match.group(2)}")
        total = Decimal("0")
        for row in value_sheet.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
            for cell in row:
                number = _decimal(cell.value)
                if number is not None:
                    total += number
        return total
    product_match = re.fullmatch(r"=([A-Z]+\d+)\*([0-9.]+)%", expression)
    if product_match:
        value = _decimal(value_sheet[product_match.group(1)].value)
        return value * Decimal(product_match.group(2)) / Decimal("100") if value is not None else None
    product_match = re.fullmatch(r"=([A-Z]+\d+)\*([0-9.]+)", expression)
    if product_match:
        value = _decimal(value_sheet[product_match.group(1)].value)
        return value * Decimal(product_match.group(2)) if value is not None else None
    return None


def map_workbook(path: str | Path) -> tuple[dict[str, Any], dict[str, Any], dict[str, Any]]:
    source = Path(path).resolve()
    formulas = load_workbook(source, data_only=False, read_only=False)
    values = load_workbook(source, data_only=True, read_only=False)
    facts: dict[str, Any] = {"schema_version": 1, "source": str(source), "sheets": []}
    trace: dict[str, Any] = {"schema_version": 1, "source": str(source), "cells": []}
    conflicts: dict[str, Any] = {"schema_version": 1, "source": str(source), "items": []}

    for raw_sheet in formulas.worksheets:
        value_sheet = values[raw_sheet.title]
        header_row, columns = _detect_header(raw_sheet)
        sheet_fact: dict[str, Any] = {
            "name": raw_sheet.title,
            "max_row": raw_sheet.max_row,
            "max_column": raw_sheet.max_column,
            "header_row": header_row,
            "columns": columns,
            "rows": [],
            "summary": [],
            "groups": [],
        }

        for row in raw_sheet.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                cached = value_sheet[cell.coordinate].value
                trace["cells"].append(
                    {
                        "sheet": raw_sheet.title,
                        "cell": cell.coordinate,
                        "row": cell.row,
                        "column": cell.column,
                        "raw": _json_value(cell.value),
                        "cached": _json_value(cached),
                    }
                )
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    calculated = _evaluate_formula(cell.value, value_sheet)
                    cached_number = _decimal(cached)
                    if calculated is None:
                        conflicts["items"].append(
                            {
                                "severity": "warning",
                                "rule": "formula_not_evaluated",
                                "sheet": raw_sheet.title,
                                "cell": cell.coordinate,
                                "formula": cell.value,
                                "cached": _json_value(cached),
                            }
                        )
                    elif cached_number is None:
                        conflicts["items"].append(
                            {
                                "severity": "warning",
                                "rule": "formula_cached_missing",
                                "sheet": raw_sheet.title,
                                "cell": cell.coordinate,
                                "formula": cell.value,
                                "calculated": str(calculated),
                                "cached": None,
                            }
                        )
                    elif abs(calculated - cached_number) > Decimal("0.01"):
                        conflicts["items"].append(
                            {
                                "severity": "error",
                                "rule": "formula_cached_mismatch",
                                "sheet": raw_sheet.title,
                                "cell": cell.coordinate,
                                "formula": cell.value,
                                "calculated": str(calculated),
                                "cached": _json_value(cached),
                            }
                        )

        if header_row:
            current_group: str | None = None
            group_provided_total = Decimal("0")
            group_recalculated_total = Decimal("0")
            group_recalculation_complete = True
            for row_number in range(header_row + 1, raw_sheet.max_row + 1):
                record: dict[str, Any] = {"row": row_number}
                for key, column in columns.items():
                    value, formula = _cell_value(raw_sheet, value_sheet, row_number, column)
                    record[key] = _json_value(value)
                    if formula:
                        record[f"{key}_formula"] = formula
                if not any(record.get(key) not in (None, "") for key in columns):
                    continue
                label = " ".join(
                    _text(raw_sheet.cell(row=row_number, column=column).value)
                    for column in range(1, raw_sheet.max_column + 1)
                )
                serial_text = str(record.get("serial") or "").strip()
                has_amount_fields = any(_decimal(record.get(key)) is not None for key in ("quantity", "unit_price", "total"))
                if not record.get("item") and not has_amount_fields and re.match(r"^[一二三四五六七八九十]+[、.]", serial_text):
                    current_group = serial_text
                    group_provided_total = Decimal("0")
                    group_recalculated_total = Decimal("0")
                    group_recalculation_complete = True
                    sheet_fact["groups"].append({"row": row_number, "name": current_group})
                    continue
                if current_group:
                    record["group"] = current_group
                is_summary = bool(re.search(r"小计|合计|税费|总计|报价有效期", label))
                if is_summary:
                    sheet_fact["summary"].append(record)
                else:
                    sheet_fact["rows"].append(record)

                quantity = _decimal(record.get("quantity"))
                unit_price = _decimal(record.get("unit_price"))
                total = _decimal(record.get("total"))
                if current_group and is_summary and re.search(r"小计|合计", label):
                    if total is not None and abs(total - group_provided_total) > Decimal("0.01"):
                        conflicts["items"].append(
                            {
                                "severity": "error",
                                "rule": "group_total_source_sum_mismatch",
                                "sheet": raw_sheet.title,
                                "row": row_number,
                                "item": current_group,
                                "provided_total": str(total),
                                "calculated_total": str(group_provided_total.quantize(Decimal("0.01"))),
                                "difference": str((total - group_provided_total).quantize(Decimal("0.01"))),
                            }
                        )
                    if total is not None and group_recalculation_complete and abs(total - group_recalculated_total) > Decimal("0.01"):
                        conflicts["items"].append(
                            {
                                "severity": "error",
                                "rule": "group_total_recalculated_mismatch",
                                "sheet": raw_sheet.title,
                                "row": row_number,
                                "item": current_group,
                                "provided_total": str(total),
                                "calculated_total": str(group_recalculated_total.quantize(Decimal("0.01"))),
                                "difference": str((total - group_recalculated_total).quantize(Decimal("0.01"))),
                            }
                        )
                    record["source_line_total_sum"] = str(group_provided_total.quantize(Decimal("0.01")))
                    record["recalculated_line_total_sum"] = (
                        str(group_recalculated_total.quantize(Decimal("0.01"))) if group_recalculation_complete else None
                    )
                    current_group = None
                elif current_group and not is_summary:
                    if total is not None:
                        group_provided_total += total
                    if quantity is not None and unit_price is not None:
                        group_recalculated_total += quantity * unit_price
                    elif total is not None:
                        group_recalculation_complete = False
                if quantity is not None and unit_price is not None and total is not None:
                    calculated = (quantity * unit_price).quantize(Decimal("0.01"))
                    difference = (total - calculated).quantize(Decimal("0.01"))
                    if abs(difference) > Decimal("0.01"):
                        conflicts["items"].append(
                            {
                                "severity": "error",
                                "rule": "quantity_unit_price_total_mismatch",
                                "sheet": raw_sheet.title,
                                "row": row_number,
                                "item": record.get("item"),
                                "quantity": str(quantity),
                                "unit_price": str(unit_price),
                                "provided_total": str(total),
                                "calculated_total": str(calculated),
                                "difference": str(difference),
                            }
                        )
        facts["sheets"].append(sheet_fact)

    conflicts["summary"] = {
        "errors": sum(item["severity"] == "error" for item in conflicts["items"]),
        "warnings": sum(item["severity"] == "warning" for item in conflicts["items"]),
    }
    return facts, trace, conflicts


def _write_json(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def _write_conflicts_markdown(path: Path, payload: dict[str, Any]) -> None:
    lines = [
        "# 项目来源数据冲突清单",
        "",
        f"错误：{payload['summary']['errors']}；警告：{payload['summary']['warnings']}。",
        "",
        "| # | 级别 | 工作表 | 单元格/行 | 项目 | 检查项 | 来源值 | 复算值 | 差异 |",
        "|---:|---|---|---|---|---|---:|---:|---:|",
    ]
    for index, item in enumerate(payload["items"], 1):
        lines.append(
            "| {index} | {severity} | {sheet} | {location} | {item} | {rule} | {provided} | {calculated} | {difference} |".format(
                index=index,
                severity=item["severity"],
                sheet=item.get("sheet", ""),
                location=item.get("cell", item.get("row", "")),
                item=item.get("item", ""),
                rule=item["rule"],
                provided=item.get("provided_total", item.get("cached", "")),
                calculated=item.get("calculated_total", item.get("calculated", "")),
                difference=item.get("difference", ""),
            )
        )
    lines.append("")
    path.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(description="项目 Excel 来源映射与算术复核")
    parser.add_argument("xlsx", help="项目 Excel 文件")
    parser.add_argument("--out-dir", required=True, help="输出 project_facts/source_trace/conflicts JSON 的目录")
    parser.add_argument("--fail-on-conflict", action="store_true", help="存在错误级算术冲突时退出 1")
    args = parser.parse_args()

    facts, trace, conflicts = map_workbook(args.xlsx)
    output = Path(args.out_dir)
    _write_json(output / "project_facts.json", facts)
    _write_json(output / "source_trace.json", trace)
    _write_json(output / "conflicts.json", conflicts)
    _write_conflicts_markdown(output / "conflicts.md", conflicts)
    print(f"OK 来源映射: {output.resolve()}")
    print(f"冲突: 错误 {conflicts['summary']['errors']} | 警告 {conflicts['summary']['warnings']}")
    for item in conflicts["items"][:20]:
        print(f"[{item['severity']}] {item['rule']} {item.get('sheet')} {item.get('cell', item.get('row', ''))}")
    if args.fail_on_conflict and conflicts["summary"]["errors"]:
        sys.exit(1)


if __name__ == "__main__":
    main()

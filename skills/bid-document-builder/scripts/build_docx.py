# -*- coding: utf-8 -*-
"""
INPUT: 中间稿 Markdown（{{RED:}} 红字、{{COVER}}/{{TOC}}/{{PAGEBREAK}} 版式标记、
       #~#### 标题、| 表格、**加粗**/- 列表/> 引用 等行内 markdown），依赖 python-docx
OUTPUT: 正式 Word 文档（.docx）：红字渲染 + 封面/目录域/章起页/页眉页脚 + 待确认清单；
        成稿后自动扫描 markdown/标记残留，错误级残留退出码 1
POS: bid-document-builder skill 的成稿引擎——保证交付文档不出现任何原始标记字符

用法:
    python build_docx.py draft.md output.docx [--ref 参考文件.doc/docx]
        [--cover-data 封面字段.json] [--cover-mode auto|reference|default] [--header 页眉文字]

标记语法:
    {{RED:内容}}              -> 红色加粗，进入待确认清单
    {{COVER}} ... {{/COVER}}  -> 封面页（段落居中，封面后自动分页）
    {{TOC}}                   -> 目录页（Word 目录域，打开后 Ctrl+A + F9 生成页码）
    {{PAGEBREAK}}             -> 强制分页
    # / ## / ### / ####       -> 标题 1/2/3/4（一级标题若前面无分页会自动另起一页）
    | a | b |                 -> 表格；|<| 左合并；|^| 上合并；单元格内 <br> 换行
    **加粗** / *斜体* / `等宽` -> 渲染为对应字体格式，符号本身不输出
    - 列表项                   -> 圆点列表（按缩进分级）
    > 引用                     -> 缩进段落
    --- （整行）               -> 忽略（分隔线不进入正文）

版式基线（2026-07-04 用户验收，勿回退）：正文宋体小四、首行缩进 2 字符；
黑体分级标题；A4；每章独立起页；表格五号字；页眉项目名 + 页脚页码域。
多级编号：章节编号一律作为文字照抄，不使用 Word 自动编号。
--ref 提供参考 .doc/.docx 时，auto 模式优先克隆参考首页（含图片、字体、行距、位置）并替换字段；
参考不可用或无参考时才使用内置 A4 中文标书样式。
"""
from __future__ import annotations

import os
import re
import sys
import argparse

sys.stdout.reconfigure(encoding="utf-8")
sys.stderr.reconfigure(encoding="utf-8")

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from cover_template import (
    clone_cover,
    collect_cover_slot_values,
    copy_first_section_page_setup,
    detect_cover_element_count,
    load_cover_values,
    prepare_reference_docx,
    replace_cover_fields,
)

RED_MARK = re.compile(r"\{\{RED:(.*?)\}\}", re.S)
COVER_BLOCK = re.compile(r"\{\{COVER\}\}(.*?)\{\{/COVER\}\}", re.S)
PAGEBREAK_LINE = re.compile(r"^\{\{PAGEBREAK\}\}\s*$")
TOC_LINE = re.compile(r"^\{\{TOC\}\}\s*$")
HEADING_LINE = re.compile(r"^(#{1,6})\s+(.*)")
BULLET_LINE = re.compile(r"^(\s*)[-*+]\s+(.+)$")
QUOTE_LINE = re.compile(r"^\s*>\s?(.*)$")
HR_LINE = re.compile(r"^(-{3,}|\*{3,}|_{3,})$")
# 行内 markdown：加粗 / 斜体（排除 5*6 类算式）/ 反引号；渲染后符号不落入成稿
# 加粗单独一层先解析（BOLD_SPAN）——加粗内可能包 {{RED:}}，若先拆红字会把 ** 拆到两段导致失配残留
BOLD_SPAN = re.compile(r"\*\*(.+?)\*\*|__(.+?)__")
INLINE_MD = re.compile(
    r"\*\*(?P<bold>.+?)\*\*"
    r"|__(?P<bold2>.+?)__"
    r"|(?<![*0-9A-Za-z])\*(?P<italic>[^*\s](?:[^*\n]*?[^*\s])?)\*(?![*0-9A-Za-z])"
    r"|`(?P<code>[^`\n]+?)`"
)
RED = RGBColor(0xFF, 0x00, 0x00)

# 成稿残留扫描：error 级必须修 draft 重出，warn 级人工判断
RESIDUE_CHECKS = [
    (re.compile(r"\{\{|\}\}"), "未识别/未闭合的 {{ }} 标记", "error"),
    (re.compile(r"\*\*"), "未闭合的 ** 加粗符", "warn"),
    (re.compile(r"```"), "``` 代码围栏残留", "warn"),
    (re.compile(r"^#{1,6}\s"), "行首 # 标题符残留", "warn"),
    (re.compile(r"^[-*+]\s"), "行首列表符残留", "warn"),
    (re.compile(r"\|[^|]+\|"), "疑似未成表的 | 竖线", "warn"),
]


def set_run_font(run, east_asia: str = "宋体", ascii_font: str = "Times New Roman", size_pt: float | None = None, bold: bool | None = None):
    run.font.name = ascii_font
    if run._element.rPr is None:
        run._element.get_or_add_rPr()
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    if bold is not None:
        run.font.bold = bold


def set_first_line_indent_chars(paragraph, chars: int = 2) -> None:
    """按"字符"设置首行缩进（firstLineChars 优先于磅值，中文文档标准做法）"""
    pPr = paragraph._p.get_or_add_pPr()
    ind = pPr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        pPr.append(ind)
    ind.set(qn("w:firstLineChars"), str(chars * 100))
    ind.set(qn("w:firstLine"), str(chars * 240))  # 兜底：按小四 12pt 折算 twips


def add_page_break(doc: Document) -> None:
    doc.add_page_break()


def setup_default_page(doc: Document) -> None:
    """A4 + 常见标书页边距"""
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.left_margin = Cm(2.8)
        section.right_margin = Cm(2.6)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)


def ensure_heading_styles(doc: Document) -> None:
    """保证 Heading 1-4 为黑体分级（母版缺失时创建）"""
    specs = [
        ("Heading 1", 16, True),
        ("Heading 2", 14, True),
        ("Heading 3", 12, True),
        ("Heading 4", 12, True),
    ]
    for name, size, bold in specs:
        try:
            style = doc.styles[name]
        except KeyError:
            style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.font.bold = bold
        style.font.size = Pt(size)
        style.font.name = "Times New Roman"
        if style.element.rPr is not None:
            style.element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")


def new_document(ref_path: str | None) -> tuple[Document, str | None]:
    prepared_ref = prepare_reference_docx(ref_path) if ref_path and os.path.exists(ref_path) else None
    if prepared_ref:
        doc = Document(prepared_ref)
        body = doc.element.body
        for child in list(body):
            if child.tag != qn("w:sectPr"):
                body.remove(child)
        ensure_heading_styles(doc)
        return doc, prepared_ref
    doc = Document()
    setup_default_page(doc)
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    style.font.size = Pt(12)  # 小四
    pf = style.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    ensure_heading_styles(doc)
    return doc, None


def strip_inline_md(text: str) -> str:
    """去掉行内 markdown 符号，保留内容（用于待确认清单/页眉/目录等纯文本场景）"""
    return INLINE_MD.sub(
        lambda m: next(g for g in (m.group("bold"), m.group("bold2"), m.group("italic"), m.group("code")) if g is not None),
        text,
    )


def plain_text(text: str) -> str:
    """去 RED 标记 + 去行内 markdown 的纯文本"""
    return strip_inline_md(RED_MARK.sub(lambda m: m.group(1), text)).strip()


def _emit_inline_runs(paragraph, text: str, size_pt: float | None, bold: bool | None = None) -> None:
    """把一段非红字文本按行内 markdown（斜体/反引号）拆成带格式的 run"""
    pos = 0
    for m in INLINE_MD.finditer(text):
        if m.start() > pos:
            run = paragraph.add_run(text[pos:m.start()])
            set_run_font(run, size_pt=size_pt, bold=bold)
        if m.group("bold") is not None or m.group("bold2") is not None:
            run = paragraph.add_run(m.group("bold") or m.group("bold2"))
            set_run_font(run, size_pt=size_pt, bold=True)
        elif m.group("italic") is not None:
            run = paragraph.add_run(m.group("italic"))
            set_run_font(run, size_pt=size_pt, bold=bold)
            run.font.italic = True
        else:
            run = paragraph.add_run(m.group("code"))
            set_run_font(run, size_pt=size_pt, bold=bold)
        pos = m.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, size_pt=size_pt, bold=bold)


def _add_red_aware(paragraph, text: str, confirm_items: list, location: str,
                   size_pt: float | None, bold: bool | None = None) -> None:
    """拆红字标记；红字恒为红色加粗，其余文本继承 bold 基调再解析斜体/反引号"""
    pos = 0
    for m in RED_MARK.finditer(text):
        if m.start() > pos:
            _emit_inline_runs(paragraph, text[pos:m.start()], size_pt, bold)
        red_text = strip_inline_md(m.group(1))
        run = paragraph.add_run(red_text)
        set_run_font(run, bold=True, size_pt=size_pt)
        run.font.color.rgb = RED
        confirm_items.append((location, red_text.strip()))
        pos = m.end()
    if pos < len(text):
        _emit_inline_runs(paragraph, text[pos:], size_pt, bold)


def add_runs(paragraph, text: str, confirm_items: list, location: str, size_pt: float | None = None) -> None:
    """三层渲染：加粗跨度 → 红字标记 → 斜体/反引号；任何标记符号都不落入成稿文本。
    加粗必须最先解析，否则 **{{RED:x}}** 的 ** 会被红字拆到两段而失配残留。"""
    pos = 0
    for m in BOLD_SPAN.finditer(text):
        if m.start() > pos:
            _add_red_aware(paragraph, text[pos:m.start()], confirm_items, location, size_pt)
        _add_red_aware(paragraph, m.group(1) or m.group(2), confirm_items, location, size_pt, bold=True)
        pos = m.end()
    if pos < len(text):
        _add_red_aware(paragraph, text[pos:], confirm_items, location, size_pt)


def is_table_row(line: str) -> bool:
    s = line.strip()
    return s.startswith("|") and s.count("|") >= 2


def is_separator_row(line: str) -> bool:
    return bool(re.fullmatch(r"\|[\s:|\-]+\|?", line.strip()))


def split_cells(line: str) -> list:
    s = line.strip().replace("\\|", "\x00").strip("|")
    return [c.strip().replace("\x00", "|") for c in s.split("|")]


def render_cover(doc: Document, cover_text: str, confirm_items: list) -> None:
    """封面：逐行居中，首行大号标题"""
    raw_lines = [ln.strip() for ln in cover_text.strip().splitlines() if ln.strip()]
    # 顶部留白
    for _ in range(3):
        doc.add_paragraph()
    for idx, line in enumerate(raw_lines):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(12 if idx else 24)
        p.paragraph_format.space_after = Pt(12)
        add_runs(p, line, confirm_items, "封面")
        size = 22 if idx == 0 else (18 if idx == 1 else 14)
        east = "黑体" if idx <= 1 else "宋体"
        for run in p.runs:
            set_run_font(run, east_asia=east, size_pt=size, bold=(idx <= 1 or run.font.bold))
    add_page_break(doc)


def render_toc(doc: Document) -> None:
    """目录页：插入 Word TOC 域（跟随 Heading 1-3），交付时提醒用户 F9 更新"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("目  录")
    set_run_font(run, east_asia="黑体", size_pt=18, bold=True)
    doc.add_paragraph()
    p = doc.add_paragraph()
    r_begin = p.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    r_begin._element.append(fld_begin)
    r_begin._element.append(instr)
    r_begin._element.append(fld_sep)
    r_result = p.add_run("（目录页码待生成：在 Word 中 Ctrl+A 全选后按 F9 更新域）")
    set_run_font(r_result, size_pt=10.5)
    r_end = p.add_run()
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    r_end._element.append(fld_end)
    add_page_break(doc)


def set_header_footer(doc: Document, header_text: str) -> None:
    """简单页眉页脚：页眉项目名，页脚页码"""
    for section in doc.sections:
        # 参考封面要求首页无页眉页脚；正文从第二页开始显示。
        section.different_first_page_header_footer = True
        header = section.header
        header.is_linked_to_previous = False
        hp = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
        hp.clear()
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = hp.add_run(header_text[:80] if header_text else "")
        set_run_font(run, east_asia="宋体", size_pt=9)

        footer = section.footer
        footer.is_linked_to_previous = False
        fp = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        fp.clear()
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        # 页码域
        run1 = fp.add_run("— ")
        set_run_font(run1, size_pt=9)
        fld_begin = OxmlElement("w:fldChar")
        fld_begin.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = "PAGE"
        fld_end = OxmlElement("w:fldChar")
        fld_end.set(qn("w:fldCharType"), "end")
        r2 = fp.add_run()
        r2._element.append(fld_begin)
        r2._element.append(instr)
        r2._element.append(fld_end)
        set_run_font(r2, size_pt=9)
        run3 = fp.add_run(" —")
        set_run_font(run3, size_pt=9)


def preprocess_blocks(raw: str) -> tuple[str, list[str]]:
    """抽出 COVER 块，正文中 COVER 替换为占位行以便顺序处理"""
    covers: list[str] = []

    def _sub(m):
        covers.append(m.group(1))
        return f"%%COVER_{len(covers) - 1}%%"

    body = COVER_BLOCK.sub(_sub, raw)
    return body, covers


def render_table(doc: Document, rows: list[list[str]], confirm_items: list, current_heading: str) -> None:
    """表格：五号字、首行加粗、</^ 合并、单元格内 <br> 换行、行首 - 转圆点"""
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    for ri, row in enumerate(rows):
        for ci in range(ncols):
            cell_text = row[ci] if ci < len(row) else ""
            if cell_text in ("<", "^"):
                continue
            cell = table.cell(ri, ci)
            para = cell.paragraphs[0]
            for run in list(para.runs):
                run._element.getparent().remove(run._element)
            # <br> 拆多段；行首列表符转圆点，避免原始字符入稿
            pieces = re.split(r"<br\s*/?>", cell_text, flags=re.I)
            for pi, piece in enumerate(pieces):
                piece = re.sub(r"^[-*+]\s+", "• ", piece.strip())
                target = para if pi == 0 else cell.add_paragraph()
                set_first_line_indent_chars(target, 0)
                add_runs(target, piece, confirm_items,
                         f"{current_heading} > 表格第{ri + 1}行", size_pt=10.5)
                if ri == 0:
                    for run in target.runs:
                        run.font.bold = True
    for ri, row in enumerate(rows):
        for ci in range(ncols):
            tok = row[ci] if ci < len(row) else ""
            try:
                if tok == "<" and ci > 0:
                    table.cell(ri, ci - 1).merge(table.cell(ri, ci))
                elif tok == "^" and ri > 0:
                    table.cell(ri - 1, ci).merge(table.cell(ri, ci))
            except Exception as e:
                print(f"警告: 表格({ri + 1},{ci + 1})合并失败({e})，已跳过", file=sys.stderr)
    for r in table.rows:
        for cell in r.cells:
            paras = cell.paragraphs
            for p in list(paras[1:]):
                if not p.text.strip():
                    p._element.getparent().remove(p._element)


def iter_doc_texts(doc: Document):
    for p in doc.paragraphs:
        yield ("正文", p.text)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    yield ("表格", p.text)


def scan_residue(doc: Document) -> tuple[list[str], list[str]]:
    """成稿后残留扫描：任何 markdown/标记原始字符都不允许悄悄交付"""
    errors, warns = [], []
    seen = set()
    for where, text in iter_doc_texts(doc):
        t = text.strip()
        if not t:
            continue
        for pat, label, level in RESIDUE_CHECKS:
            if pat.search(t):
                item = f"{label}：[{where}] {t[:40]}"
                if item in seen:
                    break
                seen.add(item)
                (errors if level == "error" else warns).append(item)
                break  # 每段只报最严重的一条
    return errors, warns


def build(
    draft_path: str,
    out_path: str,
    ref_path: str | None,
    header: str | None = None,
    cover_data_path: str | None = None,
    cover_mode: str = "auto",
) -> int:
    with open(draft_path, encoding="utf-8-sig") as f:
        raw = f.read()

    body, covers = preprocess_blocks(raw)
    lines = body.splitlines()

    doc, prepared_ref = new_document(ref_path)
    reference_cover_count = 0
    reference_cover_source: Document | None = None
    cover_values = load_cover_values(cover_data_path)
    if cover_mode in ("auto", "reference") and prepared_ref:
        reference_cover_source = Document(prepared_ref)
        reference_cover_count, boundary_method = detect_cover_element_count(reference_cover_source, prepared_ref)
        if reference_cover_count <= 0:
            if cover_mode == "reference":
                raise RuntimeError("未能识别参考文件封面边界；请改用 --cover-mode default。")
            reference_cover_source = None
        else:
            copy_first_section_page_setup(reference_cover_source, doc)
            print(f"OK 参考封面已识别: {reference_cover_count} 个元素 ({boundary_method})")
    confirm_items: list = []
    current_heading = "（文首）"
    header_guess = ""
    i = 0
    in_code = False
    content_emitted = False   # 已输出过正文内容（防止文档开头空白页）
    last_was_break = True     # 刚分过页/文档起点（防止与 {{PAGEBREAK}} 重复分页）

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # 封面占位
        m_cover = re.fullmatch(r"%%COVER_(\d+)%%", stripped)
        if m_cover:
            idx = int(m_cover.group(1))
            if idx < len(covers):
                if reference_cover_source is not None:
                    # 首页直接克隆：保留图片、字体、字号、段落间距、缩进及页面位置。
                    clone_cover(reference_cover_source, doc, reference_cover_count)
                    replace_cover_fields(doc, reference_cover_count, cover_values)
                    for slot in collect_cover_slot_values(doc, reference_cover_count):
                        if re.fullmatch(r"[\[【].+[\]】]", slot):
                            confirm_items.append(("封面", slot))
                    print("OK 已使用参考文件封面；未命中的字段保留参考占位，供终审确认")
                else:
                    render_cover(doc, covers[idx], confirm_items)
                # 页眉用封面第一行
                first = plain_text(covers[idx].strip().splitlines()[0]) if covers[idx].strip() else ""
                if first:
                    header_guess = first
            content_emitted = True
            last_was_break = True
            i += 1
            continue

        # 代码围栏：围栏行不输出，围栏内逐行原样成段（不解析标记）
        if stripped.startswith("```"):
            in_code = not in_code
            i += 1
            continue
        if in_code:
            if stripped:
                p = doc.add_paragraph()
                run = p.add_run(line.rstrip())
                set_run_font(run, size_pt=10.5)
                set_first_line_indent_chars(p, 0)
                content_emitted = True
                last_was_break = False
            i += 1
            continue

        if PAGEBREAK_LINE.match(stripped):
            if not last_was_break:
                add_page_break(doc)
            last_was_break = True
            i += 1
            continue

        if TOC_LINE.match(stripped):
            render_toc(doc)
            content_emitted = True
            last_was_break = True
            i += 1
            continue

        if is_table_row(line):
            rows = []
            while i < len(lines) and is_table_row(lines[i]):
                if not is_separator_row(lines[i]):
                    rows.append(split_cells(lines[i]))
                i += 1
            if rows:
                render_table(doc, rows, confirm_items, current_heading)
                content_emitted = True
                last_was_break = False
            continue

        m = HEADING_LINE.match(line)
        if m:
            level = min(len(m.group(1)), 4)
            title_src = re.sub(r"\s+#+\s*$", "", m.group(2).strip())
            heading_text = plain_text(title_src)
            current_heading = heading_text
            if not header_guess and heading_text:
                header_guess = heading_text
            # 章独立起页：一级标题前若作者没写 {{PAGEBREAK}}，自动补一次分页
            if level == 1 and content_emitted and not last_was_break:
                add_page_break(doc)
            fallback = False
            try:
                p = doc.add_heading("", level=level)
            except KeyError:
                fallback = True
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(12)
            add_runs(p, title_src, confirm_items, heading_text)
            if fallback:
                for run in p.runs:
                    set_run_font(run, east_asia="黑体", size_pt=18 - level * 2, bold=True)
            else:
                for run in p.runs:
                    set_run_font(run, east_asia="黑体", size_pt=None, bold=True)
            if level == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            content_emitted = True
            last_was_break = False
            i += 1
            continue

        # 分隔线：markdown 装饰行，不进入成稿
        if HR_LINE.match(stripped):
            i += 1
            continue

        # 圆点列表：按缩进分级，列表符不入稿
        m_bullet = BULLET_LINE.match(line)
        if m_bullet:
            depth = min(len(m_bullet.group(1).replace("\t", "    ")) // 2, 3)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.74 * (depth + 1))
            set_first_line_indent_chars(p, 0)
            bullet_run = p.add_run("• ")
            set_run_font(bullet_run, size_pt=12)
            add_runs(p, m_bullet.group(2), confirm_items, current_heading)
            for run in p.runs:
                if run.font.size is None:
                    set_run_font(run, size_pt=12)
            content_emitted = True
            last_was_break = False
            i += 1
            continue

        # 引用块：> 符号转缩进段落
        m_quote = QUOTE_LINE.match(line)
        if m_quote and stripped.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.74)
            set_first_line_indent_chars(p, 0)
            add_runs(p, m_quote.group(1), confirm_items, current_heading)
            for run in p.runs:
                if run.font.size is None:
                    set_run_font(run, size_pt=12)
            content_emitted = True
            last_was_break = False
            i += 1
            continue

        if stripped:
            # 正文段：首行缩进 2 字符（版式基线）；已有的全角空格缩进剥掉避免双重缩进
            text = stripped.lstrip("　")
            p = doc.add_paragraph()
            set_first_line_indent_chars(p, 2)
            add_runs(p, text, confirm_items, current_heading)
            for run in p.runs:
                if run.font.size is None:
                    set_run_font(run, size_pt=12)
            content_emitted = True
            last_was_break = False
        i += 1

    header_text = header or header_guess
    if header_text:
        try:
            set_header_footer(doc, header_text)
        except Exception as e:
            print(f"警告: 页眉页脚设置失败({e})", file=sys.stderr)

    os.makedirs(os.path.dirname(os.path.abspath(out_path)) or ".", exist_ok=True)
    doc.save(out_path)

    checklist_path = os.path.splitext(out_path)[0] + "-待确认清单.md"
    with open(checklist_path, "w", encoding="utf-8") as f:
        f.write(f"# 待确认清单（红字项共 {len(confirm_items)} 处）\n\n")
        f.write("> 终审时逐项核对，确认无误后在 Word 中将对应红字改回黑色。\n\n")
        f.write("| # | 所在位置 | 红字内容 | 确认 |\n|---|---|---|---|\n")
        for n, (loc, val) in enumerate(confirm_items, 1):
            safe = val.replace("|", "\\|")
            f.write(f"| {n} | {loc} | {safe} | ☐ |\n")

    print(f"OK 已生成: {out_path}")
    print(f"OK 待确认清单: {checklist_path}  (红字 {len(confirm_items)} 处)")

    # 残留扫描：错误级（{{ }} 标记残留）必须修 draft 重出
    errors, warns = scan_residue(doc)
    for w in warns[:10]:
        print(f"[格式警告] {w}")
    for e in errors[:10]:
        print(f"[格式错误] {e}")
    if errors:
        print(f"FAIL 成稿中有 {len(errors)} 处错误级标记残留——修正 draft 后重新生成，禁止交付本文件")
        return 1
    if warns:
        print(f"提示: {len(warns)} 处格式警告，请人工确认是否本意")
    return 0


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("draft", help="中间稿 markdown 路径")
    ap.add_argument("output", help="输出 docx 路径")
    ap.add_argument("--ref", help="参考 .doc/.docx，作为正文样式与封面母版", default=None)
    ap.add_argument("--header", help="页眉文字（缺省取封面首行/首个标题）", default=None)
    ap.add_argument("--cover-data", help="封面替换字段 JSON", default=None)
    ap.add_argument(
        "--cover-mode", choices=("auto", "reference", "default"), default="auto",
        help="auto=有参考则克隆封面；reference=必须克隆；default=忽略参考封面",
    )
    args = ap.parse_args()
    sys.exit(build(args.draft, args.output, args.ref, args.header, args.cover_data, args.cover_mode))


if __name__ == "__main__":
    main()

#!/usr/bin/env python
"""Convert OCR/screenshot-extracted PLC DB tables to standard Excel format."""

from __future__ import annotations

import argparse
from pathlib import Path
import re
import sys
from typing import Dict, List, Optional, Tuple

from openpyxl import Workbook, load_workbook

OUTPUT_COLUMNS = [
    "Tag Name",
    "Address",
    "Data Type",
    "Respect Data Type",
    "Description",
]

ADDRESS_RE = re.compile(r"^DB(\d+),(X|B|W|D)(\d+)(?:\.(\d+))?$", re.IGNORECASE)
TIA_DB_RE = re.compile(r"^%?DB(\d+)\.DB(X|B|W|D)(\d+)(?:\.(\d+))?$", re.IGNORECASE)
OFFSET_RE = re.compile(r"^(\d+)(?:\.(\d+))?$")

TYPE_KIND_RULES = {
    "boolean": "X",
    "bool": "X",
    "byte": "B",
    "char": "B",
    "usint": "B",
    "sint": "B",
    "short": "W",
    "int": "W",
    "word": "W",
    "uint": "W",
    "dint": "D",
    "dword": "D",
    "udint": "D",
    "real": "D",
    "float": "D",
}

TYPE_ALIASES = {
    "bool": "Boolean",
    "boolean": "Boolean",
    "byte": "Byte",
    "char": "Char",
    "usint": "USInt",
    "sint": "SInt",
    "short": "Short",
    "int": "Int",
    "word": "Word",
    "uint": "UInt",
    "dint": "DInt",
    "dword": "DWord",
    "udint": "UDInt",
    "real": "Real",
    "float": "Real",
}

NAME_CANDIDATES = ["Name", "Tag", "Field", "字段", "字段名", "变量名"]
TYPE_CANDIDATES = ["Type", "Data Type", "Datatype", "数据类型", "类型"]
ADDR_CANDIDATES = ["Address", "Addr", "Offset", "Logical Address", "地址", "偏移"]
COMMENT_CANDIDATES = ["Comment", "Description", "备注", "注释", "描述"]
CATEGORY_CANDIDATES = ["Category", "Group", "分类", "类别"]


def as_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def normalize_header(text: str) -> str:
    return text.strip().lower().replace("_", "").replace(" ", "")


def find_header_row(ws) -> int:
    max_scan = min(40, ws.max_row)
    for row_idx in range(1, max_scan + 1):
        values = {as_text(ws.cell(row=row_idx, column=col_idx).value) for col_idx in range(1, ws.max_column + 1)}
        if not values:
            continue
        normalized = {normalize_header(v) for v in values if v}
        if any(normalize_header(item) in normalized for item in NAME_CANDIDATES) and any(
            normalize_header(item) in normalized for item in ADDR_CANDIDATES
        ):
            return row_idx
    raise ValueError("Cannot find header row in screenshot table.")


def resolve_column(col_map: Dict[str, int], explicit: Optional[str], candidates: List[str], field_name: str) -> int:
    if explicit:
        if explicit not in col_map:
            raise ValueError(f"Column '{explicit}' not found for {field_name}.")
        return col_map[explicit]

    normalized_lookup = {normalize_header(k): v for k, v in col_map.items()}
    for item in candidates:
        key = normalize_header(item)
        if key in normalized_lookup:
            return normalized_lookup[key]

    raise ValueError(f"Cannot auto-detect column for {field_name}. Please pass explicit column name.")


def infer_kind_from_type(data_type: str) -> str:
    key = data_type.strip().lower()
    if key in TYPE_KIND_RULES:
        return TYPE_KIND_RULES[key]
    raise ValueError(f"Cannot infer address kind from data type: {data_type}")


def normalize_data_type(data_type: str) -> str:
    key = data_type.strip().lower()
    return TYPE_ALIASES.get(key, data_type.strip())


def parse_absolute_address(address: str) -> Optional[Tuple[int, str, int, Optional[int]]]:
    s = address.strip().replace(" ", "")
    match = ADDRESS_RE.match(s)
    if not match:
        match = TIA_DB_RE.match(s)
    if not match:
        return None
    db_no = int(match.group(1))
    kind = match.group(2).upper()
    byte = int(match.group(3))
    bit = int(match.group(4)) if match.group(4) is not None else None
    return db_no, kind, byte, bit


def parse_offset_value(offset_text: str) -> Tuple[int, Optional[int]]:
    s = offset_text.strip().replace(" ", "")
    match = OFFSET_RE.match(s)
    if not match:
        raise ValueError(f"Invalid offset format: {offset_text}")
    byte = int(match.group(1))
    bit = int(match.group(2)) if match.group(2) is not None else None
    return byte, bit


def build_address(raw_addr: str, data_type: str, default_db: Optional[int]) -> str:
    parsed = parse_absolute_address(raw_addr)
    if parsed is not None:
        db_no, kind, byte, bit = parsed
        if kind == "X" and bit is None:
            raise ValueError(f"X address missing bit offset: {raw_addr}")
        if kind != "X" and bit is not None:
            raise ValueError(f"Non-X address should not include bit offset: {raw_addr}")
        return f"DB{db_no},{kind}{byte}.{bit}" if kind == "X" else f"DB{db_no},{kind}{byte}"

    if default_db is None:
        raise ValueError(
            f"Address '{raw_addr}' is not absolute and --default-db is missing."
        )

    kind = infer_kind_from_type(data_type)
    byte, bit = parse_offset_value(raw_addr)

    if kind == "X":
        if bit is None:
            raise ValueError(f"Boolean offset must include bit part: {raw_addr}")
        if bit < 0 or bit > 7:
            raise ValueError(f"Bit offset out of range 0..7: {raw_addr}")
        return f"DB{default_db},X{byte}.{bit}"

    return f"DB{default_db},{kind}{byte}"


def sanitize_field(name: str) -> str:
    cleaned = re.sub(r"\s+", "", name.strip())
    cleaned = cleaned.replace("/", "_").replace("\\", "_")
    if not cleaned:
        raise ValueError("Field name is empty after sanitization.")
    return cleaned


def build_tag_name(station: int, category: str, field_name: str) -> str:
    return f"{station}.{category}.{sanitize_field(field_name)}"


def convert_rows(
    input_path: Path,
    output_path: Path,
    station: int,
    default_category: Optional[str],
    default_db: Optional[int],
    sheet_name: Optional[str],
    name_col: Optional[str],
    type_col: Optional[str],
    address_col: Optional[str],
    comment_col: Optional[str],
    category_col: Optional[str],
) -> int:
    wb = load_workbook(filename=str(input_path), data_only=True)
    ws = wb[sheet_name] if sheet_name else wb.active

    header_row = find_header_row(ws)
    col_map: Dict[str, int] = {}
    for col_idx in range(1, ws.max_column + 1):
        header = as_text(ws.cell(row=header_row, column=col_idx).value)
        if header:
            col_map[header] = col_idx

    idx_name = resolve_column(col_map, name_col, NAME_CANDIDATES, "name")
    idx_type = resolve_column(col_map, type_col, TYPE_CANDIDATES, "type")
    idx_addr = resolve_column(col_map, address_col, ADDR_CANDIDATES, "address")

    idx_comment: Optional[int] = None
    try:
        if comment_col:
            idx_comment = resolve_column(col_map, comment_col, COMMENT_CANDIDATES, "comment")
        else:
            normalized_lookup = {normalize_header(k): v for k, v in col_map.items()}
            for cand in COMMENT_CANDIDATES:
                ckey = normalize_header(cand)
                if ckey in normalized_lookup:
                    idx_comment = normalized_lookup[ckey]
                    break
    except Exception:
        idx_comment = None

    idx_category: Optional[int] = None
    if category_col:
        idx_category = resolve_column(col_map, category_col, CATEGORY_CANDIDATES, "category")
    else:
        normalized_lookup = {normalize_header(k): v for k, v in col_map.items()}
        for cand in CATEGORY_CANDIDATES:
            ckey = normalize_header(cand)
            if ckey in normalized_lookup:
                idx_category = normalized_lookup[ckey]
                break

    out_wb = Workbook()
    out_ws = out_wb.active
    out_ws.title = "Sheet1"

    for col_idx, col_name in enumerate(OUTPUT_COLUMNS, start=1):
        out_ws.cell(row=1, column=col_idx, value=col_name)

    out_row = 2
    converted = 0

    for row_idx in range(header_row + 1, ws.max_row + 1):
        raw_name = as_text(ws.cell(row=row_idx, column=idx_name).value)
        raw_type = as_text(ws.cell(row=row_idx, column=idx_type).value)
        raw_addr = as_text(ws.cell(row=row_idx, column=idx_addr).value)

        if not raw_name and not raw_addr:
            continue
        if not raw_name or not raw_type or not raw_addr:
            raise ValueError(f"Row {row_idx}: name/type/address must all be present.")

        row_category = default_category or ""
        if idx_category is not None:
            cell_category = as_text(ws.cell(row=row_idx, column=idx_category).value)
            if cell_category:
                row_category = cell_category

        if not row_category:
            raise ValueError(
                f"Row {row_idx}: category missing. Provide --category or a category column."
            )

        tag_name = build_tag_name(station, row_category, raw_name)
        data_type = normalize_data_type(raw_type)
        address = build_address(raw_addr, data_type, default_db)

        description = ""
        if idx_comment is not None:
            value = ws.cell(row=row_idx, column=idx_comment).value
            if value is not None:
                description = str(value)

        values = [tag_name, address, data_type, 1, description]
        for col_idx, value in enumerate(values, start=1):
            out_ws.cell(row=out_row, column=col_idx, value=value)

        out_row += 1
        converted += 1

    if converted == 0:
        raise ValueError("No valid rows converted from screenshot table.")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    out_wb.save(str(output_path))
    return converted


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Convert DB screenshot Excel/OCR table to standard tag Excel."
    )
    parser.add_argument("--input", required=True, help="Screenshot/OCR Excel file.")
    parser.add_argument("--output", required=True, help="Output standard Excel file.")
    parser.add_argument("--station", required=True, type=int, help="Station number, e.g. 1009.")
    parser.add_argument(
        "--category",
        help="Default category (Task/Alarm/State or custom). Required if no category column.",
    )
    parser.add_argument(
        "--default-db",
        type=int,
        help="Default DB number used when input address is only offset like 160 or 16.2.",
    )
    parser.add_argument("--sheet", help="Input sheet name. Default: active sheet.")
    parser.add_argument("--name-col", help="Explicit name column header.")
    parser.add_argument("--type-col", help="Explicit type column header.")
    parser.add_argument("--address-col", help="Explicit address/offset column header.")
    parser.add_argument("--comment-col", help="Explicit comment column header.")
    parser.add_argument("--category-col", help="Explicit category column header.")
    return parser


def main() -> int:
    parser = build_parser()
    args = parser.parse_args()

    try:
        converted = convert_rows(
            input_path=Path(args.input),
            output_path=Path(args.output),
            station=args.station,
            default_category=args.category,
            default_db=args.default_db,
            sheet_name=args.sheet,
            name_col=args.name_col,
            type_col=args.type_col,
            address_col=args.address_col,
            comment_col=args.comment_col,
            category_col=args.category_col,
        )
        print(f"Output file: {args.output}")
        print(f"Rows converted: {converted}")
        return 0
    except Exception as exc:  # pylint: disable=broad-except
        print(f"Error: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())

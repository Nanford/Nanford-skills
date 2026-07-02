#!/usr/bin/env python
"""Normalize DB screenshot tables and generate multi-station PLC tag tables."""

from __future__ import annotations

import argparse
from datetime import datetime
from pathlib import Path
import re
import sys
from typing import Dict, List, Optional, Tuple

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUTPUT_COLUMNS = [
    "Tag Name",
    "Address",
    "Data Type",
    "Respect Data Type",
    "Description",
]

ADDRESS_RE = re.compile(r"^DB(\d+),(X|B|W|D)(\d+)(?:\.(\d+))?$", re.IGNORECASE)
STATION_TAG_RE = re.compile(r"^(\d+)\.")
STRIDE_TOKEN_RE = re.compile(r"^(?:DB)?(\d+)\s*=\s*(-?\d+)$", re.IGNORECASE)

TYPE_KIND_RULES = {
    "boolean": "X",
    "bool": "X",
    "byte": "B",
    "char": "B",
    "usint": "B",
    "sint": "B",
    "short": "W",
    "int": "W",
    "word": "W",
    "uint": "W",
    "dint": "D",
    "dword": "D",
    "udint": "D",
    "real": "D",
    "float": "D",
}


def as_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def find_header_row(ws) -> int:
    max_scan = min(50, ws.max_row)
    for row_idx in range(1, max_scan + 1):
        values = {as_text(ws.cell(row=row_idx, column=col_idx).value) for col_idx in range(1, ws.max_column + 1)}
        if "Tag Name" in values and "Address" in values:
            return row_idx
    raise ValueError("Cannot find header row containing 'Tag Name' and 'Address'.")


def read_input_rows(input_path: Path, sheet_name: Optional[str]) -> List[Dict[str, object]]:
    wb = load_workbook(filename=str(input_path), data_only=True)
    ws = wb[sheet_name] if sheet_name else wb.active

    header_row = find_header_row(ws)
    col_map: Dict[str, int] = {}
    for col_idx in range(1, ws.max_column + 1):
        header = as_text(ws.cell(row=header_row, column=col_idx).value)
        if header:
            col_map[header] = col_idx

    required = ["Tag Name", "Address", "Data Type"]
    missing = [name for name in required if name not in col_map]
    if missing:
        raise ValueError(f"Missing required columns: {', '.join(missing)}")

    rows: List[Dict[str, object]] = []
    for row_idx in range(header_row + 1, ws.max_row + 1):
        tag_value = ws.cell(row=row_idx, column=col_map["Tag Name"]).value
        addr_value = ws.cell(row=row_idx, column=col_map["Address"]).value

        if tag_value is None and addr_value is None:
            continue

        tag_name = as_text(tag_value)
        address = as_text(addr_value)
        if not tag_name and not address:
            continue
        if not tag_name or not address:
            raise ValueError(f"Row {row_idx}: Tag Name and Address must both be present.")

        data_type = as_text(ws.cell(row=row_idx, column=col_map["Data Type"]).value)
        respect_value = 1
        if "Respect Data Type" in col_map:
            raw_respect = ws.cell(row=row_idx, column=col_map["Respect Data Type"]).value
            if raw_respect is not None and as_text(raw_respect):
                respect_value = raw_respect

        description_value = ""
        if "Description" in col_map:
            raw_desc = ws.cell(row=row_idx, column=col_map["Description"]).value
            if raw_desc is not None:
                description_value = str(raw_desc)

        rows.append(
            {
                "Tag Name": tag_name,
                "Address": address,
                "Data Type": data_type,
                "Respect Data Type": respect_value,
                "Description": description_value,
            }
        )

    if not rows:
        raise ValueError("No rows found after parsing the header.")

    return rows


def parse_station_selector(raw: str) -> List[int]:
    cleaned = raw.replace("，", ",").replace("；", ",").replace(";", ",")
    parts = [item.strip() for item in cleaned.split(",") if item.strip()]
    if not parts:
        raise ValueError("Station selector is empty.")

    stations = set()
    for part in parts:
        if "-" in part:
            bounds = [item.strip() for item in part.split("-", maxsplit=1)]
            if len(bounds) != 2:
                raise ValueError(f"Invalid station range: {part}")
            start = int(bounds[0])
            end = int(bounds[1])
            if end < start:
                raise ValueError(f"Invalid station range: {part}. Expect start <= end.")
            stations.update(range(start, end + 1))
        else:
            stations.add(int(part))

    return sorted(stations)


def parse_stride_map(raw: str) -> Dict[int, int]:
    cleaned = raw.replace("，", ",").replace("；", ",").replace(";", ",")
    tokens = [item.strip() for item in cleaned.split(",") if item.strip()]
    if not tokens:
        raise ValueError("Stride map is empty.")

    stride_map: Dict[int, int] = {}
    for token in tokens:
        match = STRIDE_TOKEN_RE.match(token)
        if not match:
            raise ValueError(f"Invalid stride token: {token}. Expected format like DB50=20")
        db_no = int(match.group(1))
        stride = int(match.group(2))
        if stride == 0:
            raise ValueError(f"Stride for DB{db_no} cannot be 0.")
        stride_map[db_no] = stride

    return stride_map


def parse_address(address: str) -> Tuple[int, str, int, Optional[int]]:
    normalized = address.strip().replace(" ", "")
    match = ADDRESS_RE.match(normalized)
    if not match:
        raise ValueError(f"Invalid address format: {address}")

    db_no = int(match.group(1))
    kind = match.group(2).upper()
    byte_offset = int(match.group(3))
    bit_offset = int(match.group(4)) if match.group(4) is not None else None

    if kind == "X" and bit_offset is None:
        raise ValueError(f"Bit offset is required for X address: {address}")
    if kind != "X" and bit_offset is not None:
        raise ValueError(f"Bit offset is only valid for X address: {address}")
    if bit_offset is not None and not (0 <= bit_offset <= 7):
        raise ValueError(f"Bit offset must be in 0..7: {address}")

    return db_no, kind, byte_offset, bit_offset


def format_address(db_no: int, kind: str, byte_offset: int, bit_offset: Optional[int]) -> str:
    if kind == "X":
        return f"DB{db_no},{kind}{byte_offset}.{bit_offset}"
    return f"DB{db_no},{kind}{byte_offset}"


def extract_station(tag_name: str) -> Optional[int]:
    match = STATION_TAG_RE.match(tag_name)
    if not match:
        return None
    return int(match.group(1))


def infer_base_station(rows: List[Dict[str, object]], cli_base_station: Optional[int]) -> int:
    if cli_base_station is not None:
        return cli_base_station

    for row in rows:
        station = extract_station(str(row["Tag Name"]))
        if station is not None:
            return station

    raise ValueError("Cannot infer base station from Tag Name. Use --base-station.")


def validate_type_kind(data_type: str, kind: str, row_index: int) -> None:
    expected_kind = TYPE_KIND_RULES.get(data_type.strip().lower())
    if expected_kind and expected_kind != kind:
        raise ValueError(
            f"Row {row_index}: Data Type '{data_type}' expects address kind '{expected_kind}', got '{kind}'."
        )


def split_rows_by_station(rows: List[Dict[str, object]]) -> Dict[int, List[Dict[str, object]]]:
    grouped: Dict[int, List[Dict[str, object]]] = {}
    for row in rows:
        station = extract_station(str(row["Tag Name"]))
        if station is None:
            continue
        grouped.setdefault(station, []).append(row)
    return grouped


def tag_suffix(tag_name: str) -> str:
    first_dot = tag_name.find(".")
    if first_dot == -1:
        return tag_name
    return tag_name[first_dot + 1 :]


def infer_stride_map(rows: List[Dict[str, object]]) -> Dict[int, int]:
    grouped = split_rows_by_station(rows)
    stations = sorted(grouped.keys())
    if len(stations) < 2:
        raise ValueError("Cannot infer stride: need at least 2 stations in input, or provide --stride.")

    # key -> station -> byte
    observations: Dict[Tuple[str, int, str, Optional[int]], Dict[int, int]] = {}
    for station, station_rows in grouped.items():
        for row in station_rows:
            tag = str(row["Tag Name"])
            suffix = tag_suffix(tag)
            db_no, kind, byte, bit = parse_address(str(row["Address"]))
            bit_key = bit if kind == "X" else None
            key = (suffix, db_no, kind, bit_key)
            observations.setdefault(key, {})[station] = byte

    candidates_by_db: Dict[int, List[int]] = {}
    for (_, db_no, _kind, _bit), station_to_byte in observations.items():
        sorted_pairs = sorted(station_to_byte.items(), key=lambda it: it[0])
        if len(sorted_pairs) < 2:
            continue

        base_station, base_byte = sorted_pairs[0]
        for station, byte in sorted_pairs[1:]:
            station_delta = station - base_station
            byte_delta = byte - base_byte
            if station_delta == 0:
                continue
            if byte_delta % station_delta != 0:
                continue
            stride = byte_delta // station_delta
            candidates_by_db.setdefault(db_no, []).append(stride)

    if not candidates_by_db:
        raise ValueError("Cannot infer stride from input rows. Please provide --stride.")

    stride_map: Dict[int, int] = {}
    conflicts: List[str] = []

    for db_no, values in sorted(candidates_by_db.items()):
        unique_values = sorted(set(values))
        if len(unique_values) == 1:
            stride_map[db_no] = unique_values[0]
        else:
            preview = ", ".join(str(v) for v in unique_values[:5])
            conflicts.append(f"DB{db_no}: {preview}")

    if conflicts:
        detail = "; ".join(conflicts)
        raise ValueError(f"Stride conflicts inferred from input. Please provide --stride manually. {detail}")

    return stride_map


def replace_station_prefix(tag_name: str, base_station: int, target_station: int) -> str:
    prefix = f"{base_station}."
    if tag_name.startswith(prefix):
        return f"{target_station}.{tag_name[len(prefix):]}"
    return tag_name.replace(str(base_station), str(target_station), 1)


def replace_station_in_text(text: object, base_station: int, target_station: int) -> str:
    if text is None:
        return ""
    return str(text).replace(str(base_station), str(target_station))


def category_order(tag_name: str) -> int:
    if ".Task." in tag_name:
        return 1
    if ".Alarm." in tag_name:
        return 2
    if ".State." in tag_name:
        return 3
    return 4


def station_sort_value(tag_name: str) -> int:
    station = extract_station(tag_name)
    if station is None:
        return 10**9
    return station


def build_base_rows(all_rows: List[Dict[str, object]], base_station: int) -> List[Dict[str, object]]:
    grouped = split_rows_by_station(all_rows)
    base_rows = grouped.get(base_station)
    if not base_rows:
        raise ValueError(f"Base station {base_station} not found in input rows.")
    return base_rows


def generate_rows(
    base_rows: List[Dict[str, object]],
    stations: List[int],
    base_station: int,
    stride_map: Dict[int, int],
) -> List[Dict[str, object]]:
    required_dbs = set()
    for row_index, row in enumerate(base_rows, start=1):
        db_no, kind, _, _ = parse_address(str(row["Address"]))
        validate_type_kind(str(row["Data Type"]), kind, row_index)
        required_dbs.add(db_no)

    missing_dbs = sorted(db for db in required_dbs if db not in stride_map)
    if missing_dbs:
        missing_text = ", ".join(f"DB{db}" for db in missing_dbs)
        raise ValueError(f"Missing stride for: {missing_text}")

    output_rows: List[Dict[str, object]] = []
    for target_station in stations:
        delta = target_station - base_station
        for base_row in base_rows:
            db_no, kind, base_byte, base_bit = parse_address(str(base_row["Address"]))
            stride = stride_map[db_no]
            new_byte = base_byte + delta * stride
            if new_byte < 0 or new_byte > 65535:
                raise ValueError(
                    f"Address out of range for station {target_station}: DB{db_no},{kind}{new_byte}"
                )

            output_rows.append(
                {
                    "Tag Name": replace_station_prefix(str(base_row["Tag Name"]), base_station, target_station),
                    "Address": format_address(db_no, kind, new_byte, base_bit),
                    "Data Type": base_row["Data Type"],
                    "Respect Data Type": base_row["Respect Data Type"],
                    "Description": replace_station_in_text(
                        base_row.get("Description", ""), base_station, target_station
                    ),
                }
            )

    output_rows.sort(
        key=lambda row: (
            station_sort_value(str(row["Tag Name"])),
            category_order(str(row["Tag Name"])),
            str(row["Tag Name"]),
        )
    )
    return output_rows


def insert_separator_rows(rows: List[Dict[str, object]]) -> List[Dict[str, object]]:
    separated: List[Dict[str, object]] = []
    last_station: Optional[int] = None

    for row in rows:
        station = extract_station(str(row["Tag Name"]))
        if last_station is not None and station is not None and station != last_station:
            separated.append({col: "" for col in OUTPUT_COLUMNS})
        separated.append(row)
        last_station = station

    return separated


def format_station_summary(stations: List[int]) -> str:
    if not stations:
        return ""
    is_contiguous = all(stations[i] + 1 == stations[i + 1] for i in range(len(stations) - 1))
    if is_contiguous:
        return f"{stations[0]}-{stations[-1]} (共{len(stations)}个站台)"
    return f"{','.join(str(station) for station in stations)} (共{len(stations)}个站台)"


def format_stride_summary(stride_map: Dict[int, int]) -> str:
    parts = [f"DB{db_no}={stride_map[db_no]}字节" for db_no in sorted(stride_map)]
    return ", ".join(parts)


def apply_sheet_style(ws, header_row: int) -> None:
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")

    for col_idx in range(1, len(OUTPUT_COLUMNS) + 1):
        cell = ws.cell(row=header_row, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")


def autofit_columns(ws, max_col: int) -> None:
    for col_idx in range(1, max_col + 1):
        max_len = 0
        for row_idx in range(1, ws.max_row + 1):
            value = ws.cell(row=row_idx, column=col_idx).value
            if value is None:
                continue
            max_len = max(max_len, len(str(value)))
        width = min(max(max_len + 2, 12), 64)
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def write_output(
    output_path: Path,
    rows: List[Dict[str, object]],
    stations: List[int],
    base_station: int,
    stride_map: Dict[int, int],
    stride_source: str,
) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    metadata = [
        "# PLC站台点位表",
        f"生成日期: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
        f"站台范围: {format_station_summary(stations)}",
        f"基准站台: {base_station}",
        f"DB块步长: {format_stride_summary(stride_map)}",
        f"步长来源: {stride_source}",
        f"总点位数: {len([row for row in rows if as_text(row.get('Tag Name'))])}个",
    ]

    for row_idx, text in enumerate(metadata, start=1):
        ws.cell(row=row_idx, column=1, value=text)

    header_row = 8
    for col_idx, col_name in enumerate(OUTPUT_COLUMNS, start=1):
        ws.cell(row=header_row, column=col_idx, value=col_name)

    data_start_row = header_row + 1
    for row_idx, row in enumerate(rows, start=data_start_row):
        for col_idx, col_name in enumerate(OUTPUT_COLUMNS, start=1):
            ws.cell(row=row_idx, column=col_idx, value=row.get(col_name, ""))

    apply_sheet_style(ws, header_row)
    autofit_columns(ws, len(OUTPUT_COLUMNS))
    ws.freeze_panes = f"A{data_start_row}"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(output_path))


def normalize_raw_input(
    raw_input_path: Path,
    station: int,
    category: str,
    default_db: Optional[int],
    sheet_name: Optional[str],
) -> List[Dict[str, object]]:
    """Normalize a raw screenshot/OCR Excel into standard rows in memory (no temp file)."""
    from scripts.normalize_screenshot_excel import (
        find_header_row as raw_find_header,
        resolve_column,
        build_tag_name,
        normalize_data_type,
        build_address,
        NAME_CANDIDATES,
        TYPE_CANDIDATES,
        ADDR_CANDIDATES,
        COMMENT_CANDIDATES,
        normalize_header,
    )

    wb = load_workbook(filename=str(raw_input_path), data_only=True)
    ws = wb[sheet_name] if sheet_name else wb.active

    header_row = raw_find_header(ws)
    col_map: Dict[str, int] = {}
    for col_idx in range(1, ws.max_column + 1):
        header = as_text(ws.cell(row=header_row, column=col_idx).value)
        if header:
            col_map[header] = col_idx

    idx_name = resolve_column(col_map, None, NAME_CANDIDATES, "name")
    idx_type = resolve_column(col_map, None, TYPE_CANDIDATES, "type")
    idx_addr = resolve_column(col_map, None, ADDR_CANDIDATES, "address")

    idx_comment: Optional[int] = None
    normalized_lookup = {normalize_header(k): v for k, v in col_map.items()}
    for cand in COMMENT_CANDIDATES:
        ckey = normalize_header(cand)
        if ckey in normalized_lookup:
            idx_comment = normalized_lookup[ckey]
            break

    rows: List[Dict[str, object]] = []
    for row_idx in range(header_row + 1, ws.max_row + 1):
        raw_name = as_text(ws.cell(row=row_idx, column=idx_name).value)
        raw_type = as_text(ws.cell(row=row_idx, column=idx_type).value)
        raw_addr = as_text(ws.cell(row=row_idx, column=idx_addr).value)

        if not raw_name and not raw_addr:
            continue
        if not raw_name or not raw_type or not raw_addr:
            continue

        tag_name = build_tag_name(station, category, raw_name)
        data_type = normalize_data_type(raw_type)
        address = build_address(raw_addr, data_type, default_db)

        description = ""
        if idx_comment is not None:
            value = ws.cell(row=row_idx, column=idx_comment).value
            if value is not None:
                description = str(value)

        rows.append(
            {
                "Tag Name": tag_name,
                "Address": address,
                "Data Type": data_type,
                "Respect Data Type": 1,
                "Description": description,
            }
        )

    if not rows:
        raise ValueError("No valid rows found in raw input.")

    return rows


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Generate multi-station PLC DB tag table from a standard Excel file."
    )
    parser.add_argument("--input", help="Standard Excel file path (already normalized).")
    parser.add_argument(
        "--raw-input",
        help="Raw screenshot/OCR Excel file. Will be normalized automatically. "
        "Requires --station, --category, and optionally --default-db.",
    )
    parser.add_argument(
        "--stations",
        required=True,
        help="Target stations, e.g. 1009-1022 or 1009,1011-1015",
    )
    parser.add_argument(
        "--stride",
        help="Stride map, e.g. DB50=20,DB76=2. Optional when input contains >=2 stations.",
    )
    parser.add_argument(
        "--struct-size",
        help="Struct byte size per DB as stride source, e.g. DB50=30,DB76=2. "
        "Use when you know the Struct size from a single station's DB definition.",
    )
    parser.add_argument("--output", required=True, help="Output Excel file path.")
    parser.add_argument("--base-station", type=int, help="Base station number.")
    parser.add_argument("--station", type=int, help="Station number for --raw-input mode.")
    parser.add_argument("--category", help="Category for --raw-input mode (Task/Alarm/State).")
    parser.add_argument(
        "--default-db",
        type=int,
        help="Default DB number for --raw-input mode when addresses are offsets.",
    )
    parser.add_argument("--sheet", help="Sheet name, defaults to active sheet.")
    parser.add_argument(
        "--insert-separator",
        action="store_true",
        help="Insert empty rows between stations.",
    )
    return parser


def main() -> int:
    parser = build_parser()
    args = parser.parse_args()

    try:
        if not args.input and not args.raw_input:
            parser.error("Either --input or --raw-input is required.")

        output_path = Path(args.output)

        if args.raw_input:
            if not args.station:
                parser.error("--station is required with --raw-input.")
            if not args.category:
                parser.error("--category is required with --raw-input.")

            raw_path = Path(args.raw_input)
            all_rows = normalize_raw_input(
                raw_path, args.station, args.category, args.default_db, args.sheet
            )
            print(f"Raw input normalized: {len(all_rows)} rows from {raw_path.name}")
        else:
            input_path = Path(args.input)
            all_rows = read_input_rows(input_path, args.sheet)

        base_station = infer_base_station(all_rows, args.base_station)
        base_rows = build_base_rows(all_rows, base_station)
        stations = parse_station_selector(args.stations)

        if args.stride:
            stride_map = parse_stride_map(args.stride)
            stride_source = "manual"
        elif args.struct_size:
            stride_map = parse_stride_map(args.struct_size)
            stride_source = "struct_size"
        else:
            stride_map = infer_stride_map(all_rows)
            stride_source = "auto_inferred_from_input"

        generated_rows = generate_rows(base_rows, stations, base_station, stride_map)
        if args.insert_separator:
            generated_rows = insert_separator_rows(generated_rows)

        write_output(output_path, generated_rows, stations, base_station, stride_map, stride_source)

        print(f"Output file: {output_path}")
        print(f"Stations: {format_station_summary(stations)}")
        print(f"Stride: {format_stride_summary(stride_map)}")
        print(f"Stride source: {stride_source}")
        print(f"Rows written: {len([row for row in generated_rows if as_text(row.get('Tag Name'))])}")
        return 0
    except Exception as exc:  # pylint: disable=broad-except
        print(f"Error: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())

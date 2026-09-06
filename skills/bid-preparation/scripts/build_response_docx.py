"""Build a final DOCX bid response from Markdown response files.

INPUT: ordered Markdown response files generated in bid-projects/<project>/output,
       optional page-header text (project title shown on every page).
OUTPUT: one formatted DOCX response file:
        A4 page, 宋体, 小四 body, 四号 Heading 1, 小四 Heading 2-5, black text, 1.5 line spacing,
        headings with 13pt before / 8pt after spacing, annotation text（此处附：…/【…待补充…】）
        in 五号 red italic, page header +
        page-number footer, cover page, Word TOC field, and per-section page breaks;
        plus <output>.build-report.json (source list + warnings such as missing
        images, consumed by validate_bid_package.py 的人工核查清单).
POS: Final delivery builder for the bid-preparation skill (stage 6).

Markdown layout markers (HTML comments, invisible in Markdown preview):
  <!-- cover -->              file top: render this file as the centered cover page
  <!-- toc -->                insert 目录 title + Word TOC field (update in Word to fill)
  <!-- notoc -->              next heading renders as a centered bold title outside TOC
                              (used for the three navigation tables)
  <!-- break-all-headings --> file top: every heading in this file starts a new page
                              (commercial documents: 每一章节/小节独立起页)
  <!-- pagebreak -->          force a page break before the next block

Page-break rules (mirrors the 湖北中烟 winning sample):
  * each new source file starts on a new page
  * Heading 1/2 always start a new page; Heading 3/4 too when break-all-headings is on
  * exception: a heading that directly follows its parent heading stays on the same
    page (e.g. "一、商务文件" + "（一）投标函" share a page, as in the sample)

Heading numbering is expected to be embedded in the heading text itself
(一、 / （一） / 1. / 1.1), exactly as the winning sample does; no Word auto
numbering is stacked on top, so the TOC field renders clean entries.

Images: a standalone Markdown image line `![说明](路径)` inserts the picture
centered and auto-scaled to fit the printable area (relative paths resolve
against the Markdown file's directory). Missing files degrade to a visible
placeholder paragraph instead of crashing, and are listed in the JSON result.
"""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


FONT_NAME = "宋体"
BODY_SIZE_PT = 12              # 小四
HEADING1_SIZE_PT = 14          # 四号
HEADING_OTHER_SIZE_PT = 12     # 小四，用于二至五级标题
COVER_INFO_SIZE_PT = 16        # 三号
COVER_TITLE_SIZE_PT = 36       # 小初
HEADER_FOOTER_SIZE_PT = 9      # 小五
ANNOTATION_SIZE_PT = 10.5      # 五号
HEADING_SPACE_BEFORE_PT = 13
HEADING_SPACE_AFTER_PT = 8
FIRST_LINE_INDENT_CHARS = 200  # Word 的字符单位：2 个字符 = 200
FIRST_LINE_INDENT_TWIPS = 480  # 2 × 12pt，兼容不识别 firstLineChars 的客户端
BLACK = RGBColor(0, 0, 0)
ANNOTATION_RED = RGBColor(0xFF, 0x00, 0x00)
HEADING_STYLE_BY_LEVEL = {
    1: "Heading 1",
    2: "Heading 2",
    3: "Heading 3",
    4: "Heading 4",
    5: "Heading 5",
}
HEADING_PATTERN = re.compile(r"^(#{1,5})\s+(.+)$")
MARKER_PATTERN = re.compile(r"^<!--\s*([a-z-]+)\s*-->$")
# 贪婪匹配到行尾右括号，兼容含括号的文件名（如 ISO9001_00(1).jpg）
IMAGE_PATTERN = re.compile(r"^!\[([^\]]*)\]\((.+)\)$")
# 整行加粗（**……**）：用于比五级标题更深的小节题（如 1.4.2.1.1），不进目录
BOLD_LINE_PATTERN = re.compile(r"^\*\*(.+)\*\*$")
# 行内加粗（段落/表格单元格中混排的 **……** 片段）：拆分为加粗 run 渲染
INLINE_BOLD_PATTERN = re.compile(r"\*\*(.+?)\*\*")
# 行内代码反引号、删除线等：渲染时剥掉标记，避免 Markdown 泄漏进 DOCX
INLINE_CODE_PATTERN = re.compile(r"`([^`]+)`")
STRIKETHROUGH_PATTERN = re.compile(r"~~(.+?)~~")
# 水平分割线（Markdown 语法，成品中应忽略）
HORIZONTAL_RULE_PATTERN = re.compile(r"^(-{3,}|\*{3,}|_{3,})$")
# 说明性/待补占位文字（非标书正式内容，交付前必须替换或确认）：五号红色斜体
ANNOTATION_BRACKET_PATTERN = re.compile(r"^（此处附[:：][^）]{0,60}）$|^【[^】]{1,80}】$")
ANNOTATION_KEYWORDS = ("此处附", "待补充", "待回填", "待核查", "待确认", "需人工", "说明")
# 单星号斜体 `*文字*`：渲染时剥掉星号（成对才剥，避免误伤乘号等孤立星号）
EMPHASIS_PATTERN = re.compile(r"(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)")
# 行首「**引导词：**」+ 正文：碎片化清单文风，且离开本渲染器就会漏出星号
LEAD_IN_BOLD_PATTERN = re.compile(r"^\s*(?:[-*+]\s+|\d+[.、]\s*)?\*\*[^*]{1,20}[:：]\*\*\s*\S")
# 中间稿中应避免的 Markdown 痕迹（**加粗** 为合法语法，由渲染器消化，不列入）
# 「不成对的**」改用计数判定（见 has_unpaired_bold）：正则版会把合法整行加粗误报。
MARKDOWN_LEAK_PATTERNS = (
    (re.compile(r"`"), "行内代码反引号`"),
    (re.compile(r"~~.+?~~"), "删除线~~"),
    (re.compile(r"^>\s+", re.MULTILINE), "引用块>"),
    (re.compile(r"\[([^\]]+)\]\(([^)]+)\)"), "Markdown链接语法"),
    (EMPHASIS_PATTERN, "单星号斜体*"),
)
# A4 减去左右 3.18cm / 上下 2.54cm 页边距后的可打印区域
PRINTABLE_WIDTH_CM = 21.0 - 3.18 * 2
PRINTABLE_HEIGHT_CM = 29.7 - 2.54 * 2 - 1.5  # 预留页眉页脚空间


def set_run_font(run, size_pt: float, bold: bool = False) -> None:
    run.font.name = FONT_NAME
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = BLACK


def set_annotation_font(run) -> None:
    """将待补或说明文字统一成五号红色斜体，形成可见的人工核查标记。"""
    run.font.name = FONT_NAME
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    run.font.size = Pt(ANNOTATION_SIZE_PT)
    run.font.bold = False
    run.font.italic = True
    run.font.color.rgb = ANNOTATION_RED


def is_annotation_line(text: str) -> bool:
    """整行被（此处附：…）或【…】包裹且含说明类关键词的，判定为说明性文字。"""
    if not ANNOTATION_BRACKET_PATTERN.match(text):
        return False
    return any(keyword in text for keyword in ANNOTATION_KEYWORDS)


def set_first_line_indent(paragraph) -> None:
    """设置正文首行缩进 2 字符，并保留 twips 回退值以兼容 Word/WPS。"""
    p_pr = paragraph._p.get_or_add_pPr()
    indent = p_pr.find(qn("w:ind"))
    if indent is None:
        indent = OxmlElement("w:ind")
        p_pr.append(indent)
    indent.set(qn("w:firstLineChars"), str(FIRST_LINE_INDENT_CHARS))
    indent.set(qn("w:firstLine"), str(FIRST_LINE_INDENT_TWIPS))


def configure_paragraph_format(paragraph_format, before_pt: int = 0, after_pt: int = 0) -> None:
    paragraph_format.line_spacing = 1.5
    paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    paragraph_format.space_before = Pt(before_pt)
    paragraph_format.space_after = Pt(after_pt)


def configure_style(style, size_pt: float, bold: bool, before_pt: int = 0, after_pt: int = 0) -> None:
    style.font.name = FONT_NAME
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT_NAME)
    style.font.size = Pt(size_pt)
    style.font.bold = bold
    style.font.color.rgb = BLACK
    configure_paragraph_format(style.paragraph_format, before_pt, after_pt)


def configure_document_styles(document: Document) -> None:
    configure_style(document.styles["Normal"], BODY_SIZE_PT, bold=False)
    configure_style(
        document.styles["Heading 1"], HEADING1_SIZE_PT, bold=True,
        before_pt=HEADING_SPACE_BEFORE_PT, after_pt=HEADING_SPACE_AFTER_PT,
    )
    for style_name in ("Heading 2", "Heading 3", "Heading 4", "Heading 5"):
        configure_style(
            document.styles[style_name], HEADING_OTHER_SIZE_PT, bold=True,
            before_pt=HEADING_SPACE_BEFORE_PT, after_pt=HEADING_SPACE_AFTER_PT,
        )


def configure_page_layout(document: Document) -> None:
    # A4 纵向，Word 默认页边距；python-docx 默认模板是 Letter，必须显式改
    for section in document.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(3.18)
        section.right_margin = Cm(3.18)


def set_page_break_before(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    if p_pr.find(qn("w:pageBreakBefore")) is None:
        p_pr.append(OxmlElement("w:pageBreakBefore"))


def add_field_runs(
    paragraph, instruction: str, placeholder: str, size_pt: int,
    annotation_placeholder: bool = False,
) -> None:
    """Insert a Word field (e.g. PAGE / TOC) marked dirty so Word refreshes it on open."""
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")

    for element in (begin, instr, separate):
        run = paragraph.add_run()
        set_run_font(run, size_pt)
        run._element.append(element)
    if placeholder:
        placeholder_run = paragraph.add_run(placeholder)
        if annotation_placeholder:
            set_annotation_font(placeholder_run)
        else:
            set_run_font(placeholder_run, size_pt)
    end_run = paragraph.add_run()
    set_run_font(end_run, size_pt)
    end_run._element.append(end)


def configure_header_footer(document: Document, header_text: str) -> None:
    section = document.sections[0]
    if header_text:
        header_paragraph = section.header.paragraphs[0]
        header_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        header_paragraph.paragraph_format.space_before = Pt(0)
        header_paragraph.paragraph_format.space_after = Pt(0)
        run = header_paragraph.add_run(header_text)
        set_run_font(run, HEADER_FOOTER_SIZE_PT)
    footer_paragraph = section.footer.paragraphs[0]
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_paragraph.paragraph_format.space_before = Pt(0)
    footer_paragraph.paragraph_format.space_after = Pt(0)
    add_field_runs(footer_paragraph, "PAGE", "1", HEADER_FOOTER_SIZE_PT)


def markdown_table_rows(lines: list[str], start: int) -> tuple[list[list[str]], int] | None:
    if start + 1 >= len(lines) or not lines[start].lstrip().startswith("|"):
        return None
    separator = lines[start + 1].strip()
    if not re.fullmatch(r"\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?", separator):
        return None
    rows = []
    index = start
    while index < len(lines) and lines[index].lstrip().startswith("|"):
        if index != start + 1:
            cells = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
            rows.append(cells)
        index += 1
    return rows, index


def add_inline_runs(paragraph, text: str, size_pt: int, base_bold: bool = False) -> None:
    """按行内 **加粗** 标记拆分为多个 run；顺带清洗反引号/删除线/链接等泄漏标记。"""
    text = clean_inline_markdown(text)
    position = 0
    for match in INLINE_BOLD_PATTERN.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position:match.start()])
            set_run_font(run, size_pt, bold=base_bold)
        run = paragraph.add_run(match.group(1))
        set_run_font(run, size_pt, bold=True)
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        set_run_font(run, size_pt, bold=base_bold)


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    column_count = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=column_count)
    table.style = "Table Grid"
    for row_index, row in enumerate(rows):
        for column_index in range(column_count):
            cell = table.cell(row_index, column_index)
            text = row[column_index] if column_index < len(row) else ""
            paragraph = cell.paragraphs[0]
            paragraph.style = document.styles["Normal"]
            configure_paragraph_format(paragraph.paragraph_format)
            add_inline_runs(paragraph, text, BODY_SIZE_PT, base_bold=(row_index == 0))


def add_body_paragraph(
    document: Document, text: str, page_break_before: bool = False, bold: bool = False
):
    paragraph = document.add_paragraph(style="Normal")
    configure_paragraph_format(paragraph.paragraph_format)
    # 仅常规正文缩进；整行加粗的小节题、表格、标题和占位说明均顶格。
    if not bold:
        set_first_line_indent(paragraph)
    if page_break_before:
        set_page_break_before(paragraph)
    add_inline_runs(paragraph, text.strip(), BODY_SIZE_PT, base_bold=bold)
    return paragraph


def add_annotation_paragraph(document: Document, text: str, page_break_before: bool = False):
    """说明性/待补占位段：五号红色斜体且顶格，便于交付前逐条核查。"""
    paragraph = document.add_paragraph(style="Normal")
    configure_paragraph_format(paragraph.paragraph_format)
    if page_break_before:
        set_page_break_before(paragraph)
    run = paragraph.add_run(text.strip())
    set_annotation_font(run)
    return paragraph


def add_heading(document: Document, text: str, level: int, page_break_before: bool):
    paragraph = document.add_paragraph(style=HEADING_STYLE_BY_LEVEL[level])
    if page_break_before:
        set_page_break_before(paragraph)
    run = paragraph.add_run(strip_markdown_marks(text))
    size_pt = HEADING1_SIZE_PT if level == 1 else HEADING_OTHER_SIZE_PT
    set_run_font(run, size_pt, bold=True)
    return paragraph


def add_plain_title(document: Document, text: str, page_break_before: bool) -> None:
    """导航表标题：四号加粗居中，但不用 Heading 样式，避免进入目录域。"""
    paragraph = document.add_paragraph(style="Normal")
    configure_paragraph_format(
        paragraph.paragraph_format,
        before_pt=HEADING_SPACE_BEFORE_PT,
        after_pt=HEADING_SPACE_AFTER_PT,
    )
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if page_break_before:
        set_page_break_before(paragraph)
    run = paragraph.add_run(strip_markdown_marks(text))
    set_run_font(run, HEADING1_SIZE_PT, bold=True)


def add_image(
    document: Document,
    image_path: Path,
    caption: str,
    page_break_before: bool,
    warnings: list[str],
) -> None:
    """居中插入图片并按可打印区域等比缩放；文件缺失时降级为占位段落。"""
    paragraph = document.add_paragraph(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    configure_paragraph_format(paragraph.paragraph_format, before_pt=6, after_pt=6)
    if page_break_before:
        set_page_break_before(paragraph)
    if not image_path.exists():
        warnings.append(f"图片未找到: {image_path}")
        run = paragraph.add_run(f"【图片待补充：{caption or image_path.name}】")
        set_annotation_font(run)
        return
    try:
        run = paragraph.add_run()
        picture = run.add_picture(str(image_path), width=Cm(PRINTABLE_WIDTH_CM))
        max_height = Cm(PRINTABLE_HEIGHT_CM)
        if picture.height > max_height:
            scale = max_height / picture.height
            picture.width = int(picture.width * scale)
            picture.height = max_height
    except Exception as exc:  # 图片损坏/格式不支持时不中断整份文档生成
        warnings.append(f"图片插入失败: {image_path} ({exc})")
        run = paragraph.add_run(f"【图片插入失败，请在 Word 中手动插入：{image_path.name}】")
        set_annotation_font(run)


def add_page_break_paragraph(document: Document) -> None:
    paragraph = document.add_paragraph(style="Normal")
    configure_paragraph_format(paragraph.paragraph_format)
    run = paragraph.add_run()
    set_run_font(run, BODY_SIZE_PT)
    run.add_break(WD_BREAK.PAGE)


def strip_markdown_marks(text: str) -> str:
    """去掉常见 Markdown 标记，避免封面/标题/直排文本原样漏出。

    正文段落的 **加粗** 由 add_inline_runs 单独处理；本函数用于封面、标题
    等整段直排文本，以及最终兜底清洗。
    """
    text = re.sub(r"^\s*#{1,6}\s+", "", text)
    text = INLINE_CODE_PATTERN.sub(r"\1", text)
    text = STRIKETHROUGH_PATTERN.sub(r"\1", text)
    text = text.replace("**", "").replace("__", "")
    text = re.sub(r"(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1", text)
    return text.strip()


def clean_inline_markdown(text: str) -> str:
    """清洗正文/表格单元格中除 **加粗** 以外的 Markdown 痕迹。

    单星号斜体 `*文字*` 此前未清洗，会把星号原样带进 DOCX（如 `*保密承诺：*`）。
    加粗由 add_inline_runs 拆 run 处理，故此处只剥单星号，不动 `**`。
    """
    text = INLINE_CODE_PATTERN.sub(r"\1", text)
    text = STRIKETHROUGH_PATTERN.sub(r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1", text)
    text = EMPHASIS_PATTERN.sub(r"\1", text)
    return text


def has_unpaired_bold(text: str) -> bool:
    """行内 `**` 个数为奇数即不成对，会导致加粗渲染错位、星号外漏。

    注意：整行加粗 `**小节题**` 是合法写法（两个标记，偶数），不得误报。
    """
    return text.count("**") % 2 == 1


def detect_markdown_leaks(text: str) -> list[str]:
    """检测成品正文中不应残留的 Markdown 痕迹，返回命中说明列表。"""
    hits: list[str] = []
    for pattern, label in MARKDOWN_LEAK_PATTERNS:
        if pattern.search(text):
            hits.append(label)
    if has_unpaired_bold(text):
        hits.append("不成对的**标记")
    return hits


def add_cover_page(document: Document, lines: list[str]) -> None:
    """封面：信息行三号居中；含"投标文件"的主标题小初加粗居中。"""
    for line in lines:
        text = strip_markdown_marks(line)
        if not text:
            continue
        paragraph = document.add_paragraph(style="Normal")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        compact = text.replace("　", "").replace(" ", "")
        # 主标题：以"投标文件"开头且非"XX：YY"信息项（如"投标文件""投标文件（技术部分）"）
        is_cover_title = (
            compact.startswith("投标文件") and "：" not in compact and ":" not in compact
        )
        if is_cover_title:
            configure_paragraph_format(paragraph.paragraph_format, before_pt=90, after_pt=90)
            run = paragraph.add_run(text)
            set_run_font(run, COVER_TITLE_SIZE_PT, bold=True)
        else:
            configure_paragraph_format(paragraph.paragraph_format, before_pt=6, after_pt=6)
            run = paragraph.add_run(text)
            set_run_font(run, COVER_INFO_SIZE_PT, bold=False)


def add_toc_block(document: Document, page_break_before: bool) -> None:
    add_plain_title(document, "目　录", page_break_before)
    paragraph = document.add_paragraph(style="Normal")
    configure_paragraph_format(paragraph.paragraph_format)
    add_field_runs(
        paragraph,
        'TOC \\o "1-5" \\h \\z \\u',
        "【目录域：在 Word/WPS 中全选后按 F9 或右键「更新域」生成带页码目录】",
        BODY_SIZE_PT,
        annotation_placeholder=True,
    )


class MarkdownRenderer:
    """把一组 Markdown 文件渲染进同一个 Document，维护跨文件的分页状态。"""

    def __init__(self, document: Document):
        self.document = document
        self.pending_break = False
        self.last_heading_level: int | None = None  # 上一个渲染块若是标题则记录级别
        self.rendered_any_block = False
        self.warnings: list[str] = []

    def _consume_break(self) -> bool:
        pending = self.pending_break
        self.pending_break = False
        return pending

    def _heading_needs_break(self, level: int, break_all: bool) -> bool:
        # 紧跟父级标题的子标题与父级同页（样本："一、商务文件"+"（一）投标函"同页）
        follows_parent = (
            self.last_heading_level is not None and self.last_heading_level < level
        )
        default_break = level <= 2 or break_all
        return default_break and not follows_parent

    def render_file(self, markdown_path: Path, is_first_file: bool) -> None:
        lines = markdown_path.read_text(encoding="utf-8").splitlines()
        markers = {
            MARKER_PATTERN.match(line.strip()).group(1)
            for line in lines[:5]
            if MARKER_PATTERN.match(line.strip())
        }
        if "cover" in markers:
            add_cover_page(self.document, [
                line for line in lines if not MARKER_PATTERN.match(line.strip())
            ])
            self.rendered_any_block = True
            self.last_heading_level = None
            self.pending_break = True
            return

        break_all = "break-all-headings" in markers
        if not is_first_file:
            self.pending_break = True
        notoc_next_heading = False

        index = 0
        while index < len(lines):
            line = lines[index].rstrip()
            stripped = line.strip()
            if not stripped:
                index += 1
                continue

            # Markdown 水平分割线不进入成品
            if HORIZONTAL_RULE_PATTERN.match(stripped):
                index += 1
                continue

            # 引用块 > 前缀剥掉后按正文渲染，并记入泄漏警告
            if stripped.startswith(">"):
                stripped = stripped.lstrip(">").strip()
                line = stripped
                self.warnings.append(
                    f"Markdown引用块已降级为正文（请改写成标书段落）: {markdown_path.name}"
                )
                if not stripped:
                    index += 1
                    continue

            marker_match = MARKER_PATTERN.match(stripped)
            if marker_match:
                marker = marker_match.group(1)
                if marker == "pagebreak":
                    self.pending_break = True
                elif marker == "notoc":
                    notoc_next_heading = True
                elif marker == "toc":
                    add_toc_block(self.document, self._consume_break())
                    self.rendered_any_block = True
                    self.last_heading_level = None
                # cover/break-all-headings 已在文件级处理
                index += 1
                continue

            table_result = markdown_table_rows(lines, index)
            if table_result:
                rows, index = table_result
                if self._consume_break():
                    add_page_break_paragraph(self.document)
                add_table(self.document, rows)
                self.rendered_any_block = True
                self.last_heading_level = None
                continue

            image_match = IMAGE_PATTERN.match(stripped)
            if image_match:
                caption, raw_path = image_match.group(1), image_match.group(2).strip()
                image_path = Path(raw_path)
                if not image_path.is_absolute():
                    image_path = (markdown_path.parent / image_path).resolve()
                add_image(
                    self.document,
                    image_path,
                    caption,
                    page_break_before=self._consume_break(),
                    warnings=self.warnings,
                )
                self.rendered_any_block = True
                self.last_heading_level = None
                index += 1
                continue

            heading_match = HEADING_PATTERN.match(line)
            if heading_match:
                level = len(heading_match.group(1))
                text = heading_match.group(2)
                needs_break = self._consume_break() or (
                    self.rendered_any_block and self._heading_needs_break(level, break_all)
                )
                if notoc_next_heading:
                    add_plain_title(self.document, text, needs_break)
                    notoc_next_heading = False
                    self.last_heading_level = None
                else:
                    add_heading(self.document, text, level, needs_break)
                    self.last_heading_level = level
                self.rendered_any_block = True
            else:
                # 无序列表标记剥掉后按正文渲染（写作规则要求成段论述，列表应改写为（1）（2））
                if re.match(r"^[-*+]\s+", line):
                    text = re.sub(r"^[-*+]\s+", "", line)
                    warning = (
                        f"Markdown列表标记已降级为正文（请改写成（1）（2）成段论述）: "
                        f"{markdown_path.name}"
                    )
                    if warning not in self.warnings:
                        self.warnings.append(warning)
                else:
                    text = line
                bold_match = BOLD_LINE_PATTERN.match(stripped)
                # 对即将写入的正文做泄漏检测（清洗前），便于作者回改中间稿
                for leak in detect_markdown_leaks(text):
                    warning = f"Markdown痕迹[{leak}]: {markdown_path.name}"
                    if warning not in self.warnings:
                        self.warnings.append(warning)
                if is_annotation_line(stripped):
                    add_annotation_paragraph(
                        self.document,
                        strip_markdown_marks(stripped),
                        page_break_before=self._consume_break(),
                    )
                elif bold_match:
                    add_body_paragraph(
                        self.document,
                        bold_match.group(1),
                        page_break_before=self._consume_break(),
                        bold=True,
                    )
                else:
                    add_body_paragraph(
                        self.document, text, page_break_before=self._consume_break()
                    )
                self.rendered_any_block = True
                self.last_heading_level = None
            index += 1


def build_response_docx(markdown_files: list[Path], output_path: Path, header_text: str = "") -> dict:
    document = Document()
    configure_document_styles(document)
    configure_page_layout(document)
    configure_header_footer(document, header_text)
    renderer = MarkdownRenderer(document)
    for position, markdown_file in enumerate(markdown_files):
        resolved = markdown_file.expanduser().resolve()
        if not resolved.exists():
            raise FileNotFoundError(f"响应文件不存在: {resolved}")
        renderer.render_file(resolved, is_first_file=(position == 0))
    output_path = output_path.expanduser().resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    result = {
        "output": str(output_path),
        "header": header_text,
        "source_files": [str(path) for path in markdown_files],
        "warnings": renderer.warnings,
    }
    # 构建结果落盘：图片待补等警告供 validate_bid_package.py 汇总进"人工核查清单"
    report_path = output_path.with_suffix(".build-report.json")
    try:
        report_path.write_text(
            json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8"
        )
    except OSError as exc:
        renderer.warnings.append(f"构建报告写入失败: {report_path} ({exc})")
    return result


def load_manifest(manifest_path: Path) -> tuple[list[Path], str]:
    """读取 JSON 组装清单：{"header": "...", "files": ["相对/绝对路径", ...]}。

    files 中的相对路径相对于清单文件所在目录解析，保证清单可随项目目录整体移动。
    """
    manifest_path = manifest_path.expanduser().resolve()
    try:
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise SystemExit(f"组装清单读取失败: {manifest_path} ({exc})")
    files = manifest.get("files")
    if not isinstance(files, list) or not files:
        raise SystemExit(f"组装清单缺少非空 files 列表: {manifest_path}")
    base_dir = manifest_path.parent
    resolved_files = []
    for entry in files:
        entry_path = Path(entry)
        resolved_files.append(entry_path if entry_path.is_absolute() else base_dir / entry_path)
    return resolved_files, str(manifest.get("header", ""))


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build a formatted DOCX bid response.")
    parser.add_argument("--output", type=Path, required=True, help="Final DOCX path.")
    parser.add_argument(
        "--header",
        default="",
        help="Page-header text shown on every page (usually 项目名称+投标文件).",
    )
    parser.add_argument(
        "--manifest",
        type=Path,
        help="JSON 组装清单（含 header 与按序 files），与直接列出 Markdown 文件二选一.",
    )
    parser.add_argument("markdown_files", type=Path, nargs="*", help="Ordered Markdown source files.")
    args = parser.parse_args()
    if bool(args.manifest) == bool(args.markdown_files):
        parser.error("--manifest 与 Markdown 文件列表必须二选一")
    return args


def main() -> int:
    args = parse_args()
    if args.manifest:
        markdown_files, manifest_header = load_manifest(args.manifest)
        header_text = args.header or manifest_header
    else:
        markdown_files, header_text = args.markdown_files, args.header
    result = build_response_docx(markdown_files, args.output, header_text)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
